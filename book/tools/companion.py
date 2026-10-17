#!/usr/bin/env python3
"""Build the separately authored class companion; never rewrite the original book.

Run: bash book/tools/bootstrap.sh && python book/tools/companion.py
Dependencies: the existing book bootstrap (python-docx and reportlab).
Supported source syntax: # title, ## section, plain paragraphs, and - bullets.
"""
from pathlib import Path
from html import escape
import re

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import SimpleDocTemplate, Paragraph, PageBreak

ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / 'book/companion/speaking-class.md'
PDF = ROOT / 'Master-Speaking-Class-Companion.pdf'
DOCX = ROOT / 'Master-Speaking-Class-Companion.docx'
ACCENT = '#0E6E5C'


def blocks(text):
    for block in re.split(r'\n\s*\n', text.strip()):
        lines = block.splitlines()
        if all(line.startswith('- ') for line in lines):
            yield from (('bullet', line[2:]) for line in lines)
        else:
            for line in lines:
                if line.startswith('### '):
                    yield 'subheading', line[4:]
                elif line.startswith('## '):
                    yield 'heading', line[3:]
                elif line.startswith('# '):
                    yield 'title', line[2:]
                else:
                    yield 'body', line


class PackDocTemplate(SimpleDocTemplate):
    """Give every pack a usable PDF outline as well as Word navigation."""
    def afterFlowable(self, flowable):
        if isinstance(flowable, Paragraph) and flowable.style.name == 'CompanionHeading':
            self._outline_index = getattr(self, '_outline_index', 0) + 1
            key = f'section-{self._outline_index}'
            self.canv.bookmarkPage(key)
            self.canv.addOutlineEntry(flowable.getPlainText(), key, level=0, closed=False)


def footer(canvas, doc, label='Master Speaking | Class Companion | Supplementary material'):
    canvas.saveState()
    canvas.setStrokeColor(colors.HexColor(ACCENT))
    canvas.line(42, 37, A4[0] - 42, 37)
    canvas.setFont('Companion', 8)
    canvas.setFillColor(colors.HexColor('#42565A'))
    canvas.drawString(42, 24, label)
    canvas.drawRightString(A4[0] - 42, 24, str(doc.page))
    canvas.restoreState()


def build(source=SOURCE, pdf_path=PDF, docx_path=DOCX,
          title_text='Master Speaking — Special Speaking Class Companion',
          footer_label='Master Speaking | Class Companion | Supplementary material', compact=False):
    content = list(blocks(source.read_text(encoding='utf-8')))
    fontdir = ROOT / 'book/assets/fonts'
    for name, weight in [('Companion', '400'), ('CompanionBold', '700')]:
        pdfmetrics.registerFont(TTFont(name, str(fontdir / f'source-sans-3-latin-{weight}-normal.ttf')))
    styles = getSampleStyleSheet()
    body = ParagraphStyle('CompanionBody', fontName='Companion', fontSize=10.5,
                          leading=14 if compact else 14.5, spaceAfter=4 if compact else 8, splitLongWords=False)
    heading = ParagraphStyle('CompanionHeading', parent=body, fontName='CompanionBold',
                             fontSize=17, leading=21, textColor=colors.HexColor(ACCENT),
                             spaceBefore=16, spaceAfter=12, keepWithNext=True)
    title = ParagraphStyle('CompanionTitle', parent=heading, fontSize=30, leading=35)
    bullet = ParagraphStyle('CompanionBullet', parent=body, leftIndent=12, firstLineIndent=-8,
                            spaceAfter=3 if compact else 8)
    subheading = ParagraphStyle('CompanionSubheading', parent=body, fontName='CompanionBold',
                               spaceBefore=6, spaceAfter=4, keepWithNext=True)
    story = []
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.27), Inches(11.69)
    section.top_margin = section.bottom_margin = Inches(.7)
    section.left_margin = section.right_margin = Inches(.7)
    normal = doc.styles['Normal']
    normal.font.name = 'Source Sans 3'
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(4 if compact else 8)
    if compact:
        normal.paragraph_format.line_spacing = 1.2
        doc.styles['List Bullet'].paragraph_format.space_after = Pt(3)
    for name in ['Title', 'Heading 1', 'Heading 2']:
        doc.styles[name].font.name = 'Source Sans 3'
        doc.styles[name].font.color.rgb = RGBColor.from_string(ACCENT[1:])
    doc.core_properties.title = title_text
    doc.core_properties.subject = 'New supplementary teacher and student materials'
    doc.core_properties.author = ''
    section.footer.paragraphs[0].text = footer_label
    for kind, text in content:
        # Each numbered section starts on a fresh sheet for classroom copying.
        if kind == 'heading' and re.match(r'^\d+\.', text):
            story.append(PageBreak())
            doc.add_page_break()
        style = {'title': title, 'heading': heading, 'body': body, 'bullet': bullet, 'subheading': subheading}[kind]
        story.append(Paragraph(('• ' if kind == 'bullet' else '') + escape(text), style))
        if kind == 'title':
            doc.add_heading(text, 0)
        elif kind == 'heading':
            doc.add_heading(text, 1)
        elif kind == 'subheading':
            doc.add_heading(text, 2)
        elif kind == 'bullet':
            doc.add_paragraph(text, style='List Bullet')
        else:
            doc.add_paragraph(text)
    PackDocTemplate(str(pdf_path), pagesize=A4, rightMargin=42, leftMargin=42,
                      topMargin=40, bottomMargin=49,
                      title=title_text,
                      author='', allowSplitting=True).build(
                          story, onFirstPage=lambda c, d: footer(c, d, footer_label),
                          onLaterPages=lambda c, d: footer(c, d, footer_label))
    doc.save(docx_path)
    print(f'Built {pdf_path.name} and {docx_path.name} from {source.relative_to(ROOT)}')


if __name__ == '__main__':
    build()
