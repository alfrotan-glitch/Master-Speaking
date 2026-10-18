#!/usr/bin/env python3
"""Check lesson logic across every chapter, not just text presence.

Not a factual review, answer-key certification, recording-alignment test,
pronunciation-specialist approval, or classroom trial.
"""
import json
import re
from pathlib import Path
ROOT = Path(__file__).resolve().parents[2]
book = json.loads((ROOT / 'book/artifacts/book.json').read_text())
reviews = json.loads((ROOT / 'book/master/unit-review.json').read_text())
issues = []
chapters_checked = 0
mcq_count = 0
vocab_checks = {}

def check(ok, context, message):
    if not ok:
        issues.append(f'{context}: {message}')

for unit in book['units']:
    for chapter in unit['chapters']:
        chapters_checked += 1
        context = f"U{unit['number']} Ch{chapter['number']}"
        blocks = chapter['blocks']
        mcq = False
        section = None
        for i, block in enumerate(blocks):
            kind, text = block['k'], block.get('text', '')
            if kind in ('sec', 'sub'):
                mcq = bool(re.search(r'Choose the Best Answer|Multiple Choice', text, re.I))
            if kind == 'sec':
                section = text
            if mcq and kind == 'q':
                check(i+1 < len(blocks) and blocks[i+1]['k'] == 'opt', context,
                      f'Open-response question in a multiple-choice task: {text}')
            if mcq and kind == 'opt' and i and blocks[i-1]['k'] != 'opt':
                mcq_count += 1
                check(blocks[i-1]['k'] == 'q', context, 'Options have no question stem')
                options = []
                for following in blocks[i:]:
                    if following['k'] != 'opt':
                        break
                    options.append(following['text'])
                markers = [re.match(r'^([A-Da-d])[.)]', option) for option in options]
                check(len(options) >= 2 and all(markers), context, 'Malformed choice options')
                if all(markers):
                    check([m[1].lower() for m in markers] == list('abcd')[:len(options)],
                          context, 'Choice letters are not sequential')
            if section == 'Before You Read':
                check(kind not in ('fill', 'tf'), context, 'Vocabulary/true-false item misplaced in pre-reading discussion')
            if section == 'Vocabulary Check':
                check(kind != 'tf', context, 'Vocabulary blank misclassified as true/false')
            if kind != 'reading':
                check(not re.search(r'your students|Students read each|Students evaluate|Day [234]', text),
                      context, f'Teacher-facing or undefined day reference: {text}')
            check(not re.search(r'\w+\s*→\s*/.+?/\w+\s*→', text), context, 'Fused pronunciation examples')
            if kind == 'instr':
                check('Reading Toolbox' not in text, context, 'Reference names an absent vocabulary heading')
                if i and blocks[i-1]['k'] == 'instr':
                    check(blocks[i-1].get('text') != text, context, 'Repeated consecutive instruction')
        if chapter['number'] == 1:
            names = [block.get('text') for block in blocks if block['k'] == 'sec']
            if 'Vocabulary Check' in names:
                start = next(i for i, x in enumerate(blocks) if x.get('text') == 'Vocabulary Check')
                end = next(i for i, x in enumerate(blocks) if x.get('text') == 'Before You Read')
                count = sum(x['k'] == 'fill' for x in blocks[start:end])
                vocab_checks[str(unit['number'])] = count
                expected_count = 6 if unit['number'] == 6 else 5
                check(count == expected_count, context,
                      f'Expected {expected_count} source vocabulary sentences before pre-reading')
        if chapter['number'] == 4:
            texts = [x.get('text', '') for x in blocks]
            check(any('Conclusion' in text for text in texts), context, 'Presentation outline lacks a conclusion')
            check(texts.count('Sources Used') == 1 and texts.count('Reflect and Try Again') == 1,
                  context, 'Missing or duplicated source/revision routine')
            expected = reviews[unit['number']-1]
            check('Check your progress: '+expected['check'] in texts, context, 'Missing unit-specific self-check')
            check('Put it to use: '+expected['transfer'] in texts, context, 'Missing unit-specific transfer task')
            if unit['number'] == 3:
                check(not any('difference between goals and systems' in text for text in texts),
                      context, 'Copied research outline from another unit')
            if unit['number'] == 8:
                check(texts.count('Research your topic using reliable sources.') == 1,
                      context, 'Repeated source instruction')
                examples = ['to → /tə/', 'a → /ə/', 'the → /ðə/', 'for → /fə/',
                            'of → /əv/ or /ə/', 'can → /kən/']
                check(all(texts.count(example) == 1 for example in examples), context,
                      'Weak forms not preserved as six separate examples')
                check(all(any(x['k'] == 'sub' and x.get('text') == heading for x in blocks)
                          for heading in ['The Schwa in Common Words', 'The Schwa in Academic Vocabulary',
                                          'The Schwa in Function Words']), context,
                      'Pronunciation headings wrongly classified as questions')

# Verify every Unit 4 primary-stress position against the pinned reference.
reference = json.loads((ROOT / 'book/publishing/pronunciation-reference.json').read_text())
stress_table = next(block for block in book['units'][3]['chapters'][3]['blocks']
                    if block['k'] == 'table' and block['rows'][0] == ['Word', 'Stress'])
for word, pattern in stress_table['rows'][1:]:
    chunks = pattern.split()
    marked = [i+1 for i, chunk in enumerate(chunks) if chunk.isupper()]
    target = reference['entries'][word.lower()]
    nuclei = [phone for phone in target['phones'] if phone[-1:] in ['0', '1', '2']]
    expected = next(i+1 for i, phone in enumerate(nuclei) if phone.endswith('1'))
    check(marked == [expected] and expected == target['primary_stress_syllable'],
          'U4 Ch4', f'Primary stress mismatch for {word}')
    check(''.join(chunks) == word, 'U4 Ch4', f'Stress capitalization differs between columns: {word}')

# A reading title plus an image is not enough: the text must start there too.
import pymupdf
normalize = lambda text: ' '.join(text.split())
reading_openings = 0
with pymupdf.open(ROOT / 'Master-Speaking-print-7x10.pdf') as pdf:
    for unit in book['units']:
        blocks = unit['chapters'][0]['blocks']
        title = next(x['text'] for x in blocks if x['k'] == 'readingtitle')
        first = next(x['text'] for x in blocks if x['k'] == 'reading')
        found = []
        for page in pdf:
            heading_text = normalize(' '.join(s['text'] for block in page.get_text('dict')['blocks']
                for line in block.get('lines', []) for s in line['spans'] if abs(s['size']-14.5)<0.1))
            if title in heading_text:
                found.append(page)
        check(len(found) == 1, f"U{unit['number']} Ch1", 'Reading title not uniquely rendered at its own style')
        if len(found) == 1:
            reading_openings += 1
            check(normalize(first)[:60] in normalize(found[0].get_text()), f"U{unit['number']} Ch1",
                  'Reading title/image is stranded without the opening passage')

# Verify the visually important case and IPA symbols, which ASCII-normalized
# fidelity coverage alone cannot establish.
from docx import Document
from zipfile import ZipFile
from lxml import etree
word_text = ''.join(p.text for p in Document(ROOT / 'Master-Speaking-7x10.docx').paragraphs)
word_doc = Document(ROOT / 'Master-Speaking-7x10.docx')
word_text += ''.join(cell.text for table in word_doc.tables for row in table.rows for cell in row.cells)
with ZipFile(ROOT / 'Master-Speaking.epub') as archive:
    epub_text = ''.join(etree.fromstring(archive.read(name)).xpath('string()')
                        for name in archive.namelist() if name.endswith('.xhtml'))
with pymupdf.open(ROOT / 'Master-Speaking-print-7x10.pdf') as pdf:
    pdf_text = ''.join(page.get_text() for page in pdf)
for name, text in [('Word', word_text), ('EPUB', epub_text), ('PDF', pdf_text)]:
    joined = re.sub(r'\s+', '', text)
    for word, pattern in stress_table['rows'][1:]:
        check(word in joined and re.sub(r'\s+', '', pattern) in joined,
              name, f'Visible stress capitals lost: {word}')
    for example in ['to → /tə/', 'a → /ə/', 'the → /ðə/', 'for → /fə/',
                    'of → /əv/ or /ə/', 'can → /kən/',
                    'I want tə learn English.', 'English is ə global language.']:
        check(re.sub(r'\s+', '', example) in joined, name, f'IPA/reduction model lost: {example}')

report = {'ok': not issues, 'chapters_checked': chapters_checked,
          'multiple_choice_questions_checked': mcq_count, 'vocabulary_items_by_unit': vocab_checks,
          'unit_review_tasks': len(reviews), 'reading_openings_checked': reading_openings,
          'primary_stress_words_checked': len(stress_table['rows'])-1, 'issues': issues,
          'scope': 'Lesson organization and defined regression cases; not universal exercise/factual validation.'}
(ROOT / 'book/artifacts/qa_lessons.json').write_text(json.dumps(report, indent=2)+'\n')
print(json.dumps(report, indent=2))
raise SystemExit(bool(issues))
