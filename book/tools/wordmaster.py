#!/usr/bin/env python3
"""docx.py — build the editable Word master Master-Speaking-7x10.docx from
book/master/master.xhtml (verified semantic master).

Design points (locked spec):
  * 7 x 10 in trim, mirrored book margins (inside 0.85", outside 0.6",
    top 0.75", bottom 0.8")
  * two sections: front matter (lower-roman folios, title page has no folio)
    and main body (decimal folios restart at 1)
  * genuine Word heading hierarchy via custom styles based on Heading 1-4
    (Navigation Pane + updateable TOC field "\\o 1-2")
  * structural pagination: page-break-before on Unit (H1) and Chapter (H2)
    STYLES — never empty paragraphs
  * keep-with-next on all headings, widow/orphan control on body
  * running heads: verso = STYLEREF unit title, recto = book title;
    folios are real PAGE fields
  * genuine editable tables (custom table style, repeating header rows,
    non-splitting rows) and inline editable images with alt text
"""
import os
from lxml import etree
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.section import WD_SECTION_START
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
MASTER = os.path.join(ROOT, 'book', 'master', 'master.xhtml')
ASSETS = os.path.join(ROOT, 'book', 'assets', 'images')
OUT = os.path.join(ROOT, 'Master-Speaking-7x10.docx')

NS = {'x': 'http://www.w3.org/1999/xhtml'}
XH = 'http://www.w3.org/1999/xhtml'

SERIF = 'Source Serif 4'
SANS = 'Source Sans 3'
ACCENT = RGBColor(0x0E, 0x6E, 0x5C)
INK = RGBColor(0x14, 0x18, 0x1A)
GRAY = RGBColor(0x5A, 0x6B, 0x66)

TEXT_W_IN = 7 - 0.85 - 0.6      # usable text width 5.55"

# ---------------------------------------------------------------------------
def text_of(el):
    return ' '.join(''.join(el.itertext()).split())

def cls_of(el):
    return (el.get('class') or '').split()

def sub(el, xp):
    return el.findall(xp, NS)

def gif_to_png(name):
    if name.endswith('.gif'):
        twin = name[:-4] + '.png'
        if os.path.exists(os.path.join(ASSETS, twin)):
            return twin
    return name

# ---------------------------------------------------------------------------
doc = Document()

# ---- document properties ---------------------------------------------------
cp = doc.core_properties
cp.title = 'Master Speaking'
cp.author = 'Abdul Raziq Nazari'
cp.language = 'en'

# ---- settings: mirrored margins + even/odd headers ------------------------
settings = doc.settings.element
for tag in ('w:mirrorMargins', 'w:evenAndOddHeaders'):
    e = OxmlElement(tag)
    e.set(qn('w:val'), 'true')
    settings.append(e)

# ---- styles ----------------------------------------------------------------
S = doc.styles

def pstyle(name, base=None, **kw):
    st = S.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    if base:
        st.base_style = S[base]
    st.quick_style = True
    return st

def fmt(st, font=None, size=None, bold=None, italic=None, color=None,
        align=None, before=None, after=None, line=None, keep_next=None,
        page_break=None, indent_l=None, indent_h=None, widow=True):
    pf = st.paragraph_format
    if font:
        st.font.name = font
        rpr = st.element.get_or_add_rPr()
        rf = rpr.find(qn('w:rFonts'))
        if rf is None:
            rf = OxmlElement('w:rFonts')
            rpr.append(rf)
        for a in ('w:ascii', 'w:hAnsi', 'w:cs'):
            rf.set(qn(a), font)
    if size is not None:
        st.font.size = Pt(size)
    if bold is not None:
        st.font.bold = bold
    if italic is not None:
        st.font.italic = italic
    if color is not None:
        st.font.color.rgb = color
    if align is not None:
        pf.alignment = align
    if before is not None:
        pf.space_before = Pt(before)
    if after is not None:
        pf.space_after = Pt(after)
    if line is not None:
        pf.line_spacing = line
    if keep_next is not None:
        pf.keep_with_next = keep_next
    if page_break is not None:
        pf.page_break_before = page_break
    if indent_l is not None:
        pf.left_indent = Inches(indent_l)
    if indent_h is not None:
        pf.first_line_indent = Inches(indent_h)
    pf.widow_control = widow
    return st

# front matter display styles
fmt(pstyle('MS Book Title'), SERIF, 30, bold=True, color=INK,
    align=WD_ALIGN_PARAGRAPH.CENTER, before=120, after=18, line=1.05)
fmt(pstyle('MS Author'), SANS, 15, color=GRAY,
    align=WD_ALIGN_PARAGRAPH.CENTER, before=6, after=6)
fmt(pstyle('MS Copyright Line'), SERIF, 9.5, color=INK,
    align=WD_ALIGN_PARAGRAPH.CENTER, after=5, line=1.25)
fmt(pstyle('MS FM Head', base='Heading 1'), SANS, 19, bold=True, color=INK,
    before=18, after=10, page_break=False, keep_next=True)
fmt(pstyle('MS FM Sub', base='Heading 2'), SANS, 12.5, bold=True, color=ACCENT,
    before=12, after=5, page_break=False, keep_next=True)
fmt(pstyle('MS Epigraph'), SERIF, 12.5, italic=True, color=INK,
    align=WD_ALIGN_PARAGRAPH.CENTER, before=90, after=6, line=1.3)
fmt(pstyle('MS TOC Heading'), SANS, 19, bold=True, color=INK, after=12)

# unit + chapter (structural H1 / H2)
fmt(pstyle('MS Unit Title', base='Heading 1'), SANS, 25, bold=True, color=INK,
    before=36, after=10, page_break=True, keep_next=True, line=1.05)
fmt(pstyle('MS Chapter Title', base='Heading 2'), SANS, 16, bold=True,
    color=INK, before=6, after=10, page_break=True, keep_next=True)
fmt(pstyle('MS Section', base='Heading 3'), SANS, 12, bold=True, color=ACCENT,
    before=15, after=4, keep_next=True)
fmt(pstyle('MS Subsection', base='Heading 4'), SANS, 10.5, bold=True,
    color=INK, before=11, after=3, keep_next=True)
fmt(pstyle('MS Reading Title', base='Heading 4'), SERIF, 12.5, bold=True,
    italic=False, color=INK, before=12, after=5, keep_next=True)

# body
fmt(pstyle('MS Body'), SERIF, 10.5, align=WD_ALIGN_PARAGRAPH.JUSTIFY,
    after=6, line=1.2)
fmt(pstyle('MS Instruction'), SANS, 10, italic=True, color=GRAY,
    after=5, line=1.15)
fmt(pstyle('MS Label'), SANS, 10.5, bold=True, color=INK, after=4, keep_next=True)
fmt(pstyle('MS Topic'), SANS, 10.5, bold=True, color=ACCENT, after=5)
fmt(pstyle('MS Title Line'), SERIF, 10.5, italic=True, after=4)
fmt(pstyle('MS Media Title'), SANS, 11, bold=True, color=INK, after=4,
    keep_next=True)
fmt(pstyle('MS RunIn'), SANS, 10, bold=True, color=ACCENT, before=9, after=2,
    keep_next=True)
fmt(pstyle('MS Question'), SERIF, 10.5, after=5, line=1.2,
    indent_l=0.3, indent_h=-0.3)
fmt(pstyle('MS Option'), SERIF, 10.5, after=2.5, line=1.15,
    indent_l=0.55, indent_h=-0.25)
fmt(pstyle('MS Statement'), SERIF, 10.5, after=5, line=1.2, indent_l=0.3)
fmt(pstyle('MS Reading'), SERIF, 10.5, after=5, line=1.25, indent_l=0.18)
fmt(pstyle('MS Note'), SERIF, 10, italic=True, color=GRAY, after=5,
    indent_l=0.18)
fmt(pstyle('MS Figure'), SERIF, 10.5, align=WD_ALIGN_PARAGRAPH.CENTER,
    before=8, after=10)

# left accent border for reading + note styles (style-level pBdr)
def style_left_border(name, color='0E6E5C', sz=16):
    st = S[name]
    ppr = st.element.get_or_add_pPr()
    pbdr = OxmlElement('w:pBdr')
    left = OxmlElement('w:left')
    left.set(qn('w:val'), 'single')
    left.set(qn('w:sz'), str(sz))
    left.set(qn('w:space'), '8')
    left.set(qn('w:color'), color)
    pbdr.append(left)
    ppr.append(pbdr)

style_left_border('MS Reading')
style_left_border('MS Note')

# bullets: restyle the built-in List Bullet
lb = S['List Bullet']
lb.font.name = SERIF
lb.font.size = Pt(10.5)
lb.paragraph_format.space_after = Pt(3)
lb.paragraph_format.line_spacing = 1.2

# table style: borders + header-row conditional formatting
tstyle = S.add_style('MS Table', WD_STYLE_TYPE.TABLE)
tstyle.base_style = S['Table Grid']
tblpr = tstyle.element.find(qn('w:tblPr'))
if tblpr is None:
    tblpr = OxmlElement('w:tblPr')
    # tblPr must precede any style-level paragraph props we add later
    tstyle.element.append(tblpr)
borders = OxmlElement('w:tblBorders')
for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
    b = OxmlElement(f'w:{edge}')
    b.set(qn('w:val'), 'single')
    b.set(qn('w:sz'), '4')
    b.set(qn('w:color'), 'C7D4D0')
    borders.append(b)
tblpr.append(borders)
cellmar = OxmlElement('w:tblCellMar')
for edge, w in (('top', 80), ('left', 110), ('bottom', 80), ('right', 110)):
    m = OxmlElement(f'w:{edge}')
    m.set(qn('w:w'), str(w))
    m.set(qn('w:type'), 'dxa')
    cellmar.append(m)
tblpr.append(cellmar)
# firstRow conditional: accent fill + white bold text
csp = OxmlElement('w:tblStylePr')
csp.set(qn('w:type'), 'firstRow')
rpr = OxmlElement('w:rPr')
rb = OxmlElement('w:b')
rpr.append(rb)
rc = OxmlElement('w:color')
rc.set(qn('w:val'), 'FFFFFF')
rpr.append(rc)
rf = OxmlElement('w:rFonts')
for a in ('w:ascii', 'w:hAnsi'):
    rf.set(qn(a), SANS)
rpr.append(rf)
cpr = OxmlElement('w:pPr')
csp.append(rpr)
csp.append(cpr)
tcpr = OxmlElement('w:tcPr')
shd = OxmlElement('w:shd')
shd.set(qn('w:val'), 'clear')
shd.set(qn('w:fill'), '0E6E5C')
tcpr.append(shd)
csp.append(tcpr)
tstyle.element.append(csp)

# ---------------------------------------------------------------------------
# helpers
def add_par(text='', style=None):
    p = doc.add_paragraph(style=style)
    if text:
        p.add_run(text)
    return p

def add_field(par, instr, placeholder='', bold=None, size=None):
    r = par.add_run()
    fc = OxmlElement('w:fldChar')
    fc.set(qn('w:fldCharType'), 'begin')
    r._r.append(fc)
    r2 = par.add_run()
    it = OxmlElement('w:instrText')
    it.set(qn('xml:space'), 'preserve')
    it.text = f' {instr} '
    r2._r.append(it)
    r3 = par.add_run()
    fc2 = OxmlElement('w:fldChar')
    fc2.set(qn('w:fldCharType'), 'separate')
    r3._r.append(fc2)
    r4 = par.add_run(placeholder)
    if bold is not None:
        r4.bold = bold
    if size is not None:
        r4.font.size = Pt(size)
    r5 = par.add_run()
    fc3 = OxmlElement('w:fldChar')
    fc3.set(qn('w:fldCharType'), 'end')
    r5._r.append(fc3)
    return par

def add_image(el, width_in):
    src = el.get('src') or ''
    fname = gif_to_png(os.path.basename(src))
    path = os.path.join(ASSETS, fname)
    if not os.path.exists(path):
        return
    alt = (el.get('alt') or '').strip() or 'Illustration'
    p = add_par(style='MS Figure')
    run = p.add_run()
    run.add_picture(path, width=Inches(width_in))
    # alt text on the drawing
    docpr = run._r.findall('.//' + qn('wp:docPr'))
    if docpr:
        docpr[0].set('descr', alt)
        docpr[0].set('title', alt[:80])

def qnum_reset():
    qnum_reset.n = 0

def render_list(ol_el):
    n = 0
    as_option = 'options' in cls_of(ol_el)
    for li in ol_el.findall('x:li', NS):
        n += 1
        opts = li.find('x:ol', NS)
        parts = [li.text or '']
        for c in li:
            if not isinstance(c.tag, str):
                parts.append(''.join(c.itertext()))
            elif etree.QName(c).localname != 'ol':
                parts.append(''.join(c.itertext()))
        body = ' '.join(''.join(parts).split())
        add_par(body if as_option else f'{n}. {body}',
                'MS Option' if as_option else 'MS Question')
        if opts is not None:
            for li2 in opts.findall('x:li', NS):
                add_par(' '.join(''.join(li2.itertext()).split()), 'MS Option')

def render_table(tbl_el):
    rows = tbl_el.findall('.//x:tr', NS)
    if not rows:
        return
    # split into header rows (inside thead or first row w/ th) + body
    header, body_rows = [], []
    for tr in rows:
        cells = tr.findall('x:td', NS) or tr.findall('x:th', NS)
        texts = [' '.join(''.join(c.itertext()).split()) for c in cells]
        is_head = (tr.getparent() is not None
                   and etree.QName(tr.getparent()).localname == 'thead')
        (header if is_head else body_rows).append((texts, len(cells)))
    ncols = max([len(h) for h, _ in header] + [len(r) for r, _ in body_rows] or [1])
    t = doc.add_table(rows=0, cols=ncols, style='MS Table')
    t.autofit = False
    # column widths: weighted by content with a floor (twips of 5.55" text)
    total_tw = int(TEXT_W_IN * 1440)
    weights = []
    for ci in range(ncols):
        wsum = sum(min(len(r[0][ci]) if ci < len(r[0]) else 0, 90)
                   for r in header + body_rows)
        weights.append(max(wsum, 10))
    wtotal = sum(weights)
    min_tw = int(0.75 * 1440)
    widths = [max(min_tw, int(total_tw * w / wtotal)) for w in weights]
    over = sum(widths) - total_tw
    if over > 0 and ncols > 1:
        widths[widths.index(max(widths))] -= over
    for texts, _ in header + body_rows:
        row = t.add_row()
        # rows must not split across pages
        trpr = row._tr.get_or_add_trPr()
        cs = OxmlElement('w:cantSplit')
        trpr.append(cs)
        if header and row is t.rows[0] or (len(t.rows) - 1) < len(header):
            th = OxmlElement('w:tblHeader')
            trpr.append(th)
        for ci in range(ncols):
            cell = row.cells[ci]
            cell.width = Inches(widths[ci] / 1440)
            txt = texts[ci] if ci < len(texts) else ''
            par = cell.paragraphs[0]
            par.style = S['MS Body']
            par.paragraph_format.space_after = Pt(2)
            par.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
            par.add_run(txt)
    add_par('', 'MS Body').paragraph_format.space_after = Pt(2)

def render_blocks(el, ctx=None):
    for c in el:
        if not isinstance(c.tag, str):
            continue
        ln = etree.QName(c).localname
        cls = cls_of(c)
        if ln == 'h2':
            add_par(text_of(c), 'MS FM Sub')
        elif ln == 'h3':
            add_par(text_of(c), 'MS Section')
        elif ln == 'h4' and 'reading-title' in cls:
            add_par(text_of(c), 'MS Reading Title')
        elif ln == 'h4':
            add_par(text_of(c), 'MS Subsection')
        elif ln == 'h5':
            add_par(text_of(c), 'MS RunIn')
        elif ln == 'p':
            if 'instr' in cls:
                add_par(text_of(c), 'MS Instruction')
            elif 'label' in cls:
                add_par(text_of(c), 'MS Label')
            elif 'topic' in cls:
                add_par(text_of(c), 'MS Topic')
            elif 'titleline' in cls:
                add_par(text_of(c), 'MS Title Line')
            elif 'media-title' in cls:
                add_par(text_of(c), 'MS Media Title')
            elif 'copyright-line' in cls:
                add_par(text_of(c), 'MS Copyright Line')
            elif 'author' in cls:
                add_par(text_of(c), 'MS Author')
            else:
                imgs = c.findall('.//x:img', NS)
                if imgs:
                    add_par(text_of(c), 'MS Body')
                    for im in imgs:
                        add_image(im, 3.6)
                else:
                    add_par(text_of(c), 'MS Body')
        elif ln == 'ol':
            render_list(c)
        elif ln == 'ul':
            for li in c.findall('x:li', NS):
                p = add_par(' '.join(''.join(li.itertext()).split()),
                            'List Bullet')
        elif ln == 'table':
            classes = cls
            ncols = len(c.findall('.//x:tr', NS)[0].findall('x:td', NS)) if c.findall('.//x:tr', NS) else 0
            if 'questions' in classes and ncols == 1:
                n = 0
                for td in c.findall('.//x:td', NS):
                    n += 1
                    add_par(f'{n}. ' + ' '.join(''.join(td.itertext()).split()),
                            'MS Question')
            else:
                render_table(c)
        elif ln == 'div':
            if 'reading-block' in cls:
                for cc in c:
                    if not isinstance(cc.tag, str):
                        continue
                    ln3 = etree.QName(cc).localname
                    if ln3 == 'p':
                        add_par(text_of(cc), 'MS Reading')
                    elif ln3 == 'figure':
                        for im in cc.findall('.//x:img', NS):
                            add_image(im, 4.2)
                    else:
                        add_par(text_of(cc), 'MS Reading')
            elif 'note' in cls:
                for p_el in c.findall('x:p', NS):
                    add_par(text_of(p_el), 'MS Note')
            else:
                render_blocks(c)
        elif ln == 'figure':
            for im in c.findall('.//x:img', NS):
                wide = 'unit-figure' in cls or 'reading-figure' in cls
                add_image(im, 4.6 if wide else 3.6)
        elif ln == 'blockquote':
            for p_el in c.findall('x:p', NS):
                add_par(text_of(p_el), 'MS Epigraph')

# ---------------------------------------------------------------------------
# SECTION 1 — front matter (roman folios)
sec1 = doc.sections[0]
sec1.page_width, sec1.page_height = Inches(7), Inches(10)
sec1.top_margin = Inches(0.75)
sec1.bottom_margin = Inches(0.8)
sec1.left_margin = Inches(0.85)     # inside (mirrored)
sec1.right_margin = Inches(0.6)     # outside
sec1.header_distance = Inches(0.45)
sec1.footer_distance = Inches(0.45)
sec1.different_first_page_header_footer = True   # title page: no folio
s1pr = sec1._sectPr
pgnum = OxmlElement('w:pgNumType')
pgnum.set(qn('w:fmt'), 'lowerRoman')
pgnum.set(qn('w:start'), '1')
s1pr.append(pgnum)
# front footer: centered roman folio field
fp = sec1.footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
add_field(fp, 'PAGE', 'i')

tree = etree.parse(MASTER)
body = tree.getroot().find('x:body', NS)

for part in body.findall('x:section', NS):
    kind = part.get('data-part')
    if kind == 'half-title':
        continue
    if kind == 'title-page':
        for el in part:
            if not isinstance(el.tag, str):
                continue
            ln2 = etree.QName(el).localname
            if ln2 == 'h1':
                add_par(text_of(el), 'MS Book Title')
            elif ln2 == 'p':
                add_par(text_of(el), 'MS Author')
    elif kind == 'copyright':
        for el in part.findall('x:p', NS):
            add_par(text_of(el), 'MS Copyright Line')
    elif kind == 'foreword':
        add_pagebreak = None
        head = part.find('x:h1', NS)
        add_par(text_of(head), 'MS FM Head')
        for el in part:
            if isinstance(el.tag, str) and etree.QName(el).localname == 'p':
                add_par(text_of(el), 'MS Body')
    elif kind == 'how-to-use':
        head = part.find('x:h1', NS)
        add_par(text_of(head), 'MS FM Head')
        render_blocks(part)
    elif kind == 'epigraph':
        render_blocks(part)

# TOC page (end of section 1)
tocp = add_par('Contents', 'MS TOC Heading')
tocp.paragraph_format.page_break_before = True
p = doc.add_paragraph()
add_field(p, 'TOC \\o "1-2" \\h \\z \\u',
          'Table of contents — right-click and choose “Update Field” to '
          'refresh page numbers.')

# ---------------------------------------------------------------------------
# SECTION 2 — main body (decimal folios restart at 1)
sec2 = doc.add_section(WD_SECTION_START.NEW_PAGE)
sec2.page_width, sec2.page_height = Inches(7), Inches(10)
sec2.top_margin = Inches(0.75)
sec2.bottom_margin = Inches(0.8)
sec2.left_margin = Inches(0.85)
sec2.right_margin = Inches(0.6)
sec2.header_distance = Inches(0.45)
sec2.footer_distance = Inches(0.45)
s2pr = sec2._sectPr
pgnum2 = OxmlElement('w:pgNumType')
pgnum2.set(qn('w:start'), '1')
s2pr.append(pgnum2)

# recto (default/odd) running head: book title, right-aligned
hp = sec2.header.paragraphs[0]
hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
hr = hp.add_run('MASTER SPEAKING')
hr.font.name = SANS
hr.font.size = Pt(8.5)
hr.font.color.rgb = GRAY
# verso (even) running head: STYLEREF unit title, left-aligned
ep = sec2.even_page_header.paragraphs[0]
ep.alignment = WD_ALIGN_PARAGRAPH.LEFT
er = ep.add_run()
er.font.name = SANS
er.font.size = Pt(8.5)
er.font.color.rgb = GRAY
add_field(ep, 'STYLEREF "MS Unit Title" \\* MERGEFORMAT', 'Unit', )
for r in ep.runs:
    r.font.name = SANS
    r.font.size = Pt(8.5)
    r.font.color.rgb = GRAY
# footers: centered PAGE fields (odd + even)
for f in (sec2.footer, sec2.even_page_footer):
    fp2 = f.paragraphs[0]
    fp2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_field(fp2, 'PAGE', '1')

for unit in body.findall('x:section[@class="unit"]', NS):
    unum = unit.get('data-unit')
    uname = text_of(unit.find('.//x:span[@class="unit-name"]', NS))
    add_par(f'Unit {unum} · {uname}', 'MS Unit Title')
    fig = unit.find('x:figure', NS)
    if fig is not None:
        for im in fig.findall('.//x:img', NS):
            add_image(im, 4.9)
    warm = unit.find('x:div[@class="warmup"]', NS)
    if warm is not None:
        add_par('THINK', 'MS Label')
        render_blocks(warm)
    for ch in unit.findall('x:section[@class="chapter"]', NS):
        cnum = ch.get('data-chapter')
        cname = text_of(ch.find('.//x:span[@class="chapter-name"]', NS))
        add_par(f'Chapter {cnum} · {cname}', 'MS Chapter Title')
        render_blocks(ch)

doc.save(OUT)
print(f'DOCX written: {OUT} ({os.path.getsize(OUT):,} bytes)')
