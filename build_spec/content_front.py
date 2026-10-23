# -*- coding: utf-8 -*-
"""Front matter: title, introduction & objectives, how to use (students & teachers),
contents, diagnostic test. Spec sections 2, 9, 10, 11."""
from style import (para, run, box, table, spacer, h_banner, page_break,
                   setup_doc, NAVY, TEAL, GOLD, INK, GREY, WHITE,
                   FILL_BOX_TEAL, FILL_BOX_GOLD, FILL_BOX_GREY, FILL_NAVY, FILL_TEAL)
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

PARTS = [
    ("Part One", "Build Strong Sentences", [1, 2]),
    ("Part Two", "Express Ideas Clearly", [3, 4]),
    ("Part Three", "Think Beyond the Present", [5, 6]),
    ("Part Four", "Compare and Persuade", [7, 8]),
    ("Part Five", "Describe and Explain", [9, 10]),
    ("Part Six", "Speak with Greater Precision", [11, 12]),
    ("Part Seven", "Sound More Natural and Advanced", [13, 14]),
    ("Part Eight", "Speak with Confidence", [15, 16]),
]

LESSON_TITLES = {
    1: "Tenses for Real Communication",
    2: "Telling Stories and Describing Experiences",
    3: "Questions, Auxiliary Verbs, and Follow-up Questions",
    4: "Modal Verbs for Real-Life Situations",
    5: "Conditionals for Decisions and Possibilities",
    6: "Wishes, Regrets, and Hypothetical Situations",
    7: "Comparisons Beyond the Basics",
    8: "Linking Ideas and Building Arguments",
    9: "Passive Voice in Natural English",
    10: "Causative Structures",
    11: "Gerunds and Infinitives",
    12: "Relative Clauses and Concise Descriptions",
    13: "Reported Speech and Reporting Verbs",
    14: "Participles and Advanced Sentence Patterns",
    15: "Hedging, Emphasis, and Natural Inversion",
    16: "Grammar in Real Conversation",
}


def title_page(doc):
    spacer(doc, 60)
    para(doc, "TOEFL HOUSE", size=14, bold=True, color=TEAL, align='c', space_after=2)
    para(doc, "GRAMMAR FOR", size=30, bold=True, color=NAVY, align='c', space_before=18, space_after=0)
    para(doc, "SPECIAL SPEAKING", size=30, bold=True, color=NAVY, align='c', space_before=0, space_after=6)
    para(doc, "A practical grammar course for accurate, natural, confident speech",
         size=13, italic=True, color=GREY, align='c', space_before=6, space_after=24)
    box(doc, None, [
        ("Learn the grammar. Use it in conversation. Make it part of your speech.", True),
    ], fill=FILL_BOX_TEAL, label_color=TEAL)
    spacer(doc, 40)
    para(doc, "Speak clearly. Connect naturally. Grow with every conversation.", size=12, color=INK, align='c')
    para(doc, "16 lessons  ·  3 review tests  ·  final assessment  ·  5 appendices",
         size=11, color=GREY, align='c')
    page_break(doc)


def introduction(doc):
    h_banner(doc, "INTRODUCTION", sub="What this book does and how it works")
    para(doc, "This book has one purpose: to help you use grammar accurately, naturally and "
              "confidently when you speak English.")
    para(doc, "It is not enough to know a rule. The goal is to turn that knowledge into active "
              "speaking skill — to choose the right structure, in the right moment, and say it "
              "without thinking about it.")
    para(doc, "Every lesson follows the same path: you speak first, you learn clear grammar patterns, "
              "you see how similar structures differ, you practise, and then you speak again — "
              "this time with the new grammar in your voice. Each lesson ends with a short review "
              "and a self-check, so you always know what you can already do.")
    para(doc, "Everyone works through the same lesson and speaking goal. Begin with a clear, "
              "accurate response, then experiment with more precise and flexible ways to express "
              "the same idea. Optional challenges help you grow when you are ready. "
              "Natural, appropriate language matters more than complexity for its own sake.")
    para(doc, "", space_after=6)
    para(doc, "By the end of the course, you should be able to:", size=11, bold=True, color=NAVY)
    for b in [
        "construct grammatically accurate sentences while speaking;",
        "choose appropriate tenses and sentence structures for your meaning;",
        "ask questions and respond with greater precision;",
        "explain experiences, opinions and hypothetical situations;",
        "connect ideas into coherent spoken responses;",
        "use a wider range of structures naturally;",
        "recognise and correct your recurring grammatical mistakes;",
        "apply learned structures in unfamiliar speaking situations.",
    ]:
        para(doc, f"•  {b}", indent=0.12)
    para(doc, "Completion of the sixteen lessons is not proof of these outcomes by itself. The "
              "written Diagnostic Test and Final Assessment sample related grammar skills, but they "
              "are not parallel tests and their raw totals should not be compared as a measure of "
              "gain. For a fairer progress check, use the same speaking task conditions and rubric "
              "before and after the course, and compare the four criterion scores.",
         size=11, italic=True, color=GREY)
    page_break(doc)


def publication_info(doc):
    h_banner(doc, "COPYRIGHT AND PUBLICATION INFORMATION",
             sub="Grammar for Special Speaking — TOEFL HOUSE")
    para(doc, "© TOEFL HOUSE, 2026. All rights reserved.", size=11, bold=True, color=NAVY)
    para(doc, "Grammar for Special Speaking — a practical grammar course for accurate, natural, "
              "confident speech. Author and publisher: TOEFL HOUSE. Publication year: 2026.")
    para(doc, "No part of this publication may be reproduced, stored or transmitted in any form "
              "without prior written permission from the publisher, except for brief quotations "
              "used in teaching, review and private study.")
    para(doc, "", space_after=4)
    table(doc, ["Publication detail", "Information"], [
        ["Title", "Grammar for Special Speaking"],
        ["Author / creator", "TOEFL HOUSE"],
        ["Edition", "[Edition information to be confirmed before print release]"],
        ["ISBN", "[ISBN to be assigned]"],
        ["Place of publication", "[To be confirmed]"],
        ["Companion volume", "Grammar for Special Speaking — Teacher's Book (TOEFL HOUSE)"],
        ["Language", "English"],
    ], widths=[2.0, 4.3], font_size=10)
    para(doc, "", space_after=4)
    para(doc, "Disclaimer", size=11, bold=True, color=NAVY)
    para(doc, "TOEFL® is a registered trademark of Educational Testing Service (ETS). This book is "
              "an independent teaching resource prepared by TOEFL HOUSE; it is not affiliated with, "
              "authorised or endorsed by ETS. Any mention of tests or exams in this book is for "
              "general language-teaching purposes only.")
    para(doc, "Statement of limits", size=11, bold=True, color=NAVY)
    para(doc, "This manuscript has been produced and checked with automated structural and "
              "editorial quality gates. Qualified teacher review, classroom piloting, EPUBCheck "
              "certification and a commercial print proof are separate steps recorded in the "
              "project documentation; none of them is claimed here.")
    para(doc, "The answer key for learners appears in Appendix 5. Teachers should use the "
              "companion Teacher's Book, which contains lesson procedures, full answers with "
              "notes, model speaking responses, timing and differentiation guidance.",
         size=10.5, italic=True, color=GREY)
    page_break(doc)


def how_to_use(doc):
    h_banner(doc, "HOW TO USE THIS BOOK", sub="For students and teachers")
    para(doc, "For students", size=12, bold=True, color=NAVY)
    para(doc, "Work through the lessons in order: each one builds on earlier material. Always "
              "begin with the Speaking Hook — say your answers out loud before you read the rule. "
              "The rule then gives you the tool you were already reaching for.")
    para(doc, "Do the exercises out loud whenever possible. Grammar becomes speech through the "
              "mouth, not only through the eyes. Use the Pair Speaking and Real Talk sections as "
              "they are designed: no reading from notes.")
    para(doc, "After every lesson, cover the page and try the retrieval questions. Tick the "
              "self-check honestly. Return to the Formula Book (Appendix 1) whenever you review.")
    para(doc, "", space_after=4)
    para(doc, "Practise, then stretch", size=12, bold=True, color=NAVY)
    table(doc, ["Step", "Focus", "What to do"], [
        [("Build the pattern", True),
         "accurate, understandable sentences and useful follow-up questions",
         "complete the core practice and say your answers aloud"],
        [("Stretch your answer", True),
         "more precise meanings, natural choices and connected responses",
         "try an optional challenge; explain why you chose that expression"],
    ], widths=[1.7, 2.2, 2.4])
    para(doc, "Start with a clear opinion and reasons. Then, if you can, add a qualification, "
              "respond to a different view, or give a specific example. "
              "You can revisit either step whenever it helps your speaking.")
    para(doc, "", space_after=4)
    para(doc, "For teachers", size=12, bold=True, color=NAVY)
    para(doc, "Keep grammar explanation short and give learners substantial time to produce spoken "
              "English. The suggested pacing below is a guide — adjust it to your class length.")
    table(doc, ["Activity", "Minutes (of a 60-minute lesson)"], [
        ["Explanation and examples", "15"],
        ["Controlled practice", "10"],
        ["Pair speaking", "20"],
        ["Independent speaking challenge", "10"],
        ["Feedback and review", "5"],
        [("Total", True), ("60", True)],
    ], widths=[3.6, 2.7])
    para(doc, "Review earlier structures in the review tests and in retrieval practice. Evaluate "
              "speaking with the success criteria and the rubric in the Final Assessment — not with "
              "rigid model answers. Accept alternative answers where an exercise allows them.")
    para(doc, "The companion Teacher's Book follows the same sixteen lessons and gives teachers "
              "step-by-step procedures for all eight lesson functions, complete answers with "
              "explanations, model speaking responses, realistic timings, differentiation and "
              "remediation guidance, and common-error notes.", space_before=6)
    page_break(doc)


def contents(doc):
    h_banner(doc, "CONTENTS", sub="Sixteen lessons in eight thematic parts")
    para(doc, "This is an automatic Word table of contents. It updates to show page numbers when opened "
              "in Word; if it has not refreshed, select it and choose Update Field / Update entire table.",
         size=10.5, italic=True, color=GREY)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = 4
    p.paragraph_format.space_after = 4
    begin_run = p.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    begin_run._r.append(fld_begin)
    instr_run = p.add_run()
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = ' TOC \\o "1-1" \\h \\z \\u '
    instr_run._r.append(instr)
    sep_run = p.add_run()
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    sep_run._r.append(separate)
    result = p.add_run("Update this field to populate the contents and page numbers.")
    result.font.italic = True
    result.font.color.rgb = GREY
    end_run = p.add_run()
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    end_run._r.append(fld_end)
    page_break(doc)


def diagnostic(doc, keys_registry):
    h_banner(doc, "DIAGNOSTIC TEST", sub="Find your starting point — 25 questions · no time limit")
    para(doc, "Do this test before Lesson 1. Do not study first: the point is to see what you can "
              "already do. Mark your score out of 25 and keep the paper. This is a broad starting "
              "snapshot, not a parallel form of the 30-item Final Assessment; do not interpret the "
              "difference between their raw totals as a learning gain. For a comparable progress "
              "check, complete the speaking sample below and score it with the same rubric used at "
              "the end of the book.", size=11, italic=True, color=GREY)
    para(doc, "Part A — Tenses and time  (10 questions)", size=11, bold=True, color=NAVY)
    A = [
        "I __________ (see) that film last week.",
        "She __________ (work) here since 2021.",
        "While I __________ (cook — ongoing background), the phone rang.",
        "They __________ (work) there for ten years before they moved to Kabul.",
        "At sea level, pure water __________ (boil) at 100°C.",
        "They __________ (build) a new bridge when we visited.",
        "I __________ my keys — I can't get in.  (lose)",
        "Last summer we __________ (go) to Herat.",
        "How many books __________ you __________ (read) so far this month?",
        "When I arrived, the train __________ already __________ (leave).",
    ]
    for i, q in enumerate(A, 1):
        para(doc, f"{i}.  {q}", indent=0.12)
    para(doc, "Part B — Structures  (10 questions)", size=11, bold=True, color=NAVY)
    B = [
        "If it __________ (rain), we'll stay home.",
        "I wish I __________ (know) the answer.",
        "My phone needs __________.  (repair)",
        "The report __________ (write) by the team last week.",
        "He suggested __________ early.  (leave)",
        "Could you tell me where __________ ?  (the station / be)",
        "The girl __________ helped me was very kind.  (who / that / whose / which)",
        "I'd rather you __________ (not, smoke) here.",
        "Ali is __________ than his brother.  (intelligent)",
        "Not until midnight __________.  (she / finish)",
    ]
    for i, q in enumerate(B, 11):
        para(doc, f"{i}.  {q}", indent=0.12)
    para(doc, "Part C — Use and communication  (5 questions)", size=11, bold=True, color=NAVY)
    C = [
        "I __________ find my passport anywhere.  (mustn't / can't)",
        "__________ being very tired, he finished the race.  (Despite / Although)",
        "She asked me, \"Where do you live?\"  →  She asked me __________.",
        "It __________ rain later — take an umbrella.  (may / might / could)",
        "I __________ my hair cut yesterday.  (had / got)",
    ]
    for i, q in enumerate(C, 21):
        para(doc, f"{i}.  {q}", indent=0.12)
    page_break(doc)
    para(doc, "Part D — Baseline speaking sample (not included in the 25-point grammar score)",
         size=11, bold=True, color=NAVY)
    box(doc, "Speak for two minutes", [
        "Choose one topic. Prepare for two minutes, then speak for two minutes without reading a script:",
        "•  A decision that changed your life — and what you would do differently today.",
        "•  One change that would improve education or work in your country.",
        "•  A person who influenced you, and the advice they gave you.",
        "After your talk, a partner asks two follow-up questions. Answer them, then ask one relevant "
        "question in return. If you get stuck, repair and continue. Record the topic and, if possible, "
        "save a recording. At the end of the course, repeat this task under the same conditions. The "
        "Final Assessment uses the same task and four-criterion rubric; score this baseline now and "
        "score the final sample later."
    ], fill=FILL_BOX_TEAL, label_color=TEAL)
    table(doc, ["Criterion", "4 — Strong", "3 — Good", "2 — Developing", "1 — Beginning"], [
        [("Grammatical accuracy", True), "structures are consistently accurate; occasional slips do not affect meaning", "generally accurate; some errors, meaning always clear", "frequent errors; meaning sometimes unclear", "errors dominate; meaning often unclear"],
        [("Appropriate structure selection", True), "wide range used naturally (tenses, conditionals, linkers, hedging…) for the intended meaning", "good range; choices usually match the meaning", "limited range; some structures forced or inappropriate", "basic structures only; frequent wrong choices"],
        [("Clarity", True), "ideas well organised and easy to follow; effective linking", "clear overall; some linking", "ideas loosely connected; hard to follow at times", "ideas listed without clear connection"],
        [("Successful communication", True), "message fully delivered; natural delivery, relevant follow-up exchange and self-repair", "message delivered; answers and asks relevant follow-ups, with some hesitation", "message partly delivered; responses are limited or pauses are long", "message largely not delivered; little response to questions"],
    ], widths=[1.35, 1.35, 1.25, 1.2, 1.15], font_size=8.5)
    para(doc, "Baseline speaking score: ____ / 16. Keep this page for the matched end-of-course comparison.",
         size=10.5, italic=True, color=GREY)
    para(doc, "When you finish, check your work in Appendix 5. The answer key is kept separate so you can complete the diagnostic without seeing it.",
         size=10.5, italic=True, color=GREY)
    keys_registry["Diagnostic"] = {
        "A": ["saw", "has worked / has been working", "was cooking", "had worked / had been working; 'worked' is also possible with before", "boils", "were building",
              "have lost", "went", "have read", "had already left"],
        "B": ["rains", "knew", "repairing / to be repaired", "was written", "leaving",
              "the station is", "who / that", "didn't smoke", "more intelligent", "did she finish"],
        "C": ["can't", "Despite", "where I lived / where I live (if still true)", "may / might / could", "had / got"],
    }
    para(doc, "Score guide: 21–25 — review the specific items you missed, then try the stretch "
              "tasks.  15–20 — revisit the patterns you found difficult and practise them aloud.  "
              "Below 15 — repeat the core practice in short sessions; accuracy first, speed later.  "
              "Use this score to choose what to practise next; speaking ability involves more than one test.",
         size=10.5, italic=True, color=GREY)
    page_break(doc)
