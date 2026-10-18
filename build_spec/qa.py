# -*- coding: utf-8 -*-
"""Automated QA gates for 'Grammar for Special Speaking'.
Run AFTER build.py:  python build_spec/qa.py
Exit code 0 = all gates pass. Non-zero = genuine defects (do not weaken checks)."""
import sys, os, re
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from docx import Document
from docx.oxml.ns import qn

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DOCX = os.path.join(ROOT, 'Grammar for Special Speaking (TOEFL House).docx')

SPEC_TITLES = {
    1: "Tenses for Real Communication", 2: "Telling Stories and Describing Experiences",
    3: "Questions, Auxiliary Verbs, and Follow-up Questions", 4: "Modal Verbs for Real-Life Situations",
    5: "Conditionals for Decisions and Possibilities", 6: "Wishes, Regrets, and Hypothetical Situations",
    7: "Comparisons Beyond the Basics", 8: "Linking Ideas and Building Arguments",
    9: "Passive Voice in Natural English", 10: "Causative Structures",
    11: "Gerunds and Infinitives", 12: "Relative Clauses and Concise Descriptions",
    13: "Reported Speech and Reporting Verbs", 14: "Participles and Advanced Sentence Patterns",
    15: "Hedging, Emphasis, and Natural Inversion", 16: "Grammar in Real Conversation",
}
PARTS_LABELS = ["The Speaking Hook", "The Simple Rule", "See the Difference",
                "Practice the Pattern", "Pair Speaking", "The Advanced Upgrade",
                "Real Talk Challenge", "Review and Self-Check"]


def main():
    if not os.path.exists(DOCX):
        print("FAIL: manuscript missing:", DOCX)
        return 1
    doc = Document(DOCX)
    texts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for r in t.rows:
            for c in r.cells:
                texts.append(c.text)
    blob = "\n".join(texts)
    fails = []

    def check(name, cond):
        print(("  PASS " if cond else "  FAIL ") + name)
        if not cond:
            fails.append(name)

    print("=== QA GATES: Grammar for Special Speaking ===")
    # structure
    check("16 lessons", sum(1 for t in texts if re.match(r"^LESSON \d+", t.strip())) == 16)
    for i, lab in enumerate(PARTS_LABELS, 1):
        check(f"framework Part {i} ({lab}) x16",
              sum(1 for t in texts if t.strip() == f"Part {i}  ·  {lab}") == 16)
    check("16 lesson titles match spec",
          all(f"LESSON {i}  ·  {t.upper()}" in blob for i, t in SPEC_TITLES.items()))
    # document order
    order = []
    for t in texts:
        m = re.match(r"^LESSON (\d+)", t.strip())
        if m:
            order.append(f"L{m.group(1)}")
        m = re.match(r"^REVIEW TEST (\d+)", t.strip())
        if m:
            order.append(f"T{m.group(1)}")
    expected = ["L1", "L2", "L3", "L4", "T1", "L5", "L6", "L7", "L8", "T2",
                "L9", "L10", "L11", "L12", "T3", "L13", "L14", "L15", "L16"]
    check("document order (tests in scope positions)", order == expected)
    # keys
    check("16 quick-answer boxes", sum(1 for t in texts if t.strip().startswith("Quick answers")) == 16)
    check("appendix answer key covers 16 lessons",
          sum(1 for t in texts if re.match(r"^Lesson \d+$", t.strip())) >= 16)
    # assessments
    check("3 review tests", all(f"REVIEW TEST {i}" in blob for i in (1, 2, 3)))
    check("diagnostic present", "DIAGNOSTIC TEST" in blob)
    check("final assessment + rubric + error plan",
          all(k in blob for k in ("FINAL ASSESSMENT", "Evaluation rubric", "improvement plan")))
    check("5 appendices", all(f"APPENDIX {i}" in blob for i in (1, 2, 3, 4, 5)))
    # final assessment covers all 16 lessons (keyword evidence)
    coverage = {
        "L1": "__________ (live) here since 2020", "L2": "(already / have) dinner",
        "L3": "When does the shop close", "L4": "(mustn't / can't)",
        "L5": "Third conditional: 'He didn't revise, so he failed.'",
        "L6": "wish I __________ (not/say) that yesterday",
        "L7": "far + give a reason frame", "L8": "being expensive, the course is worth it",
        "L9": "stolen the documents", "L10": "I paid someone to fix the roof",
        "L11": "look forward to __________ (meet) you", "L12": "The book what I read was amazing",
        "L13": "She advised", "L14": "Because he had studied hard, he passed easily",
        "L15": "Rarely had she felt so confident", "L16": "Speaking task (2 minutes to prepare",
    }
    for k, kw in coverage.items():
        check(f"final assessment covers {k}", kw in blob)
    # review-test scope evidence (no leakage of later-lesson grammar)
    def walk_between(start, end):
        """Collect all text (document order) strictly between two banner markers."""
        collecting = False
        out = []
        for child in doc.element.body:
            tag = child.tag.split('}')[1]
            if tag not in ('p', 'tbl'):
                continue
            t = ''.join(n.text or '' for n in child.iter(qn('w:t')))
            if not collecting and start in t:
                collecting = True
                continue
            if collecting and end in t:
                break
            if collecting:
                out.append(t)
        return '\n'.join(out)

    t1 = walk_between("REVIEW TEST 1", "LESSON 5")
    t2 = walk_between("REVIEW TEST 2", "LESSON 9")
    t3 = walk_between("REVIEW TEST 3", "LESSON 13")
    leak = [("T1", t1, ["reported", "She asked me", "participle", "hedg", "cleft", "inversion",
                        "Never have", "Not only", "causative", "have my", "gerund", "avoid"]),
            ("T2", t2, ["participle", "Having", "hedg", "may well", "cleft", "inversion",
                        "Never have", "Not only", "reported", "promised to", "which were",
                        "avoid", "look forward to", "causative", "have my", "was written",
                        "Not until"]),
            ("T3", t3, ["hedg", "may well", "it could be argued", "inversion", "Never have",
                        "Not only", "Having finished", "Not having", "reported", "She asked me",
                        "denied", "cleft"])]
    for name, scope, banned in leak:
        hits = [w for w in banned if w.lower() in scope.lower()]
        check(f"{name} scope clean (no untaught grammar) {hits if hits else ''}", not hits)
    # banned language
    check("no scheduling words", not re.search(r"\b(Day \d|Week \d|MONTH \d|teaching days)\b", blob))
    check("no Master Speaking references", not re.search(r"Master Speaking", blob, re.I))
    check("no unearned claims", not re.search(r"expert-verified|publication-ready|pedagogically validated", blob, re.I))
    check("zero decorative emojis", sum(1 for ch in blob if ord(ch) >= 0x1F300) == 0)
    # known-defect regressions
    check("no 'Although being' error pattern", "Although being" not in blob)
    check("guiding principle present", "Learn the grammar. Use it in conversation." in blob)

    print()
    if fails:
        print(f"RESULT: {len(fails)} FAILURES -> {fails}")
        return 1
    print("RESULT: ALL QA GATES PASS")
    return 0


if __name__ == "__main__":
    sys.exit(main())
