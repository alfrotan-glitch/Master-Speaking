# -*- coding: utf-8 -*-
"""Front matter for the Grammar for Special Speaking Teacher's Book."""
from style import (para, box, table, spacer, h_banner, page_break,
                   NAVY, TEAL, GOLD, INK, GREY,
                   FILL_BOX_TEAL, FILL_BOX_GOLD, FILL_BOX_GREY)
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def title_page(doc):
    spacer(doc, 60)
    para(doc, "TOEFL HOUSE", size=14, bold=True, color=TEAL, align='c', space_after=2)
    para(doc, "GRAMMAR FOR", size=28, bold=True, color=NAVY, align='c',
         space_before=18, space_after=0)
    para(doc, "SPECIAL SPEAKING", size=28, bold=True, color=NAVY, align='c',
         space_before=0, space_after=4)
    para(doc, "TEACHER'S BOOK", size=22, bold=True, color=GOLD, align='c',
         space_before=6, space_after=10)
    para(doc, "Lesson procedures, complete answers, model speaking responses, "
              "timing and differentiation for all sixteen lessons",
         size=13, italic=True, color=GREY, align='c', space_before=6, space_after=24)
    box(doc, None, [
        ("Companion to the Student Book: Grammar for Special Speaking", True),
        ("Learn the grammar. Use it in conversation. Make it part of your speech.", False),
    ], fill=FILL_BOX_TEAL, label_color=TEAL)
    spacer(doc, 40)
    para(doc, "For teachers of Intermediate and Advanced learners", size=12, color=INK, align='c')
    para(doc, "16 lesson guides  ·  3 review-test guides  ·  assessment keys  ·  feedback guidance",
         size=11, color=GREY, align='c')
    page_break(doc)


def publication_note(doc):
    h_banner(doc, "COPYRIGHT AND PUBLICATION INFORMATION",
             sub="Teacher's Book — Grammar for Special Speaking")
    para(doc, "© TOEFL HOUSE, 2026. All rights reserved.", size=11, bold=True, color=NAVY)
    para(doc, "This Teacher's Book is the companion volume to the Student Book "
              "Grammar for Special Speaking (author and publisher: TOEFL HOUSE; publication year "
              "2026). Full publication details, edition and ISBN placeholders, and the statement "
              "of limits are printed on the Student Book's copyright page and apply equally here.")
    para(doc, "The lesson guides in this volume reproduce the Student Book's objectives, practice "
              "items, answers and model responses for teaching use. Assessment keys are provided "
              "here for instructors; they should not be distributed to learners before a test is "
              "completed.")
    para(doc, "TOEFL® is a registered trademark of Educational Testing Service (ETS). This book is "
              "an independent teaching resource prepared by TOEFL HOUSE; it is not affiliated with, "
              "authorised or endorsed by ETS.")
    page_break(doc)


def introduction(doc):
    h_banner(doc, "INTRODUCTION", sub="What this Teacher's Book is and how it was built")
    para(doc, "This volume exists so that a competent English teacher can prepare and deliver any "
              "lesson in Grammar for Special Speaking without having to invent missing "
              "explanations, answers, procedures or activities.")
    para(doc, "Every guide that follows corresponds, line by line, to the Student Book: the same "
              "sixteen lessons in the same order, the same eight lesson functions, the same "
              "practice items and the same answer key. The teaching content in this book and the "
              "practice content in the Student Book are generated from a single shared source, so "
              "the two volumes cannot drift apart between editions.")
    para(doc, "Each guide gives you: measurable lesson outcomes and success criteria; materials and "
              "preparation; a sixty-minute timing plan that matches the Student Book's pacing "
              "guidance; step-by-step procedures for all eight lesson functions; quick "
              "checks for understanding; the complete practice answers with notes on where "
              "learners lose marks; model exchanges and model responses for the speaking tasks; "
              "assessment guidance for open-ended work; Intermediate scaffolding and Advanced "
              "extension; common learner errors with correction strategies; a forty-five-minute "
              "shorter-session plan; and review and feedback guidance.")
    para(doc, "The Student Book's assessment policy carries over unchanged: model answers are "
              "examples, not scripts; acceptable alternatives are accepted wherever an exercise "
              "allows them; speaking tasks are judged with success criteria and the four-criterion "
              "rubric, not with a single mandatory response; and the Diagnostic Test and Final "
              "Assessment raw totals must not be compared as a measure of learning gain.")
    para(doc, "What this book does not claim: it is not a record of classroom trials, teacher "
              "certification, or external approval. Its procedures are editorially prepared and "
              "structurally verified; qualified-teacher review and a real learner pilot remain "
              "separate, documented steps.", size=10.5, italic=True, color=GREY)
    page_break(doc)


def how_to_use(doc):
    h_banner(doc, "HOW TO USE THIS TEACHER'S BOOK", sub="Map of the guides")
    para(doc, "Lesson-by-lesson layout", size=12, bold=True, color=NAVY)
    table(doc, ["Section of each guide", "What it gives you"], [
        ["Lesson outcomes and measurable success criteria",
         "what to teach toward, and observable evidence that learners got there"],
        ["Materials and preparation",
         "what to print, write on the board or set up before learners arrive"],
        ["Suggested timing (60-minute lesson)",
         "minutes per lesson function; every table totals exactly 60"],
        ["Parts 1–8 (the eight lesson functions)",
         "step-by-step procedures mirroring the Student Book, in teaching order"],
        ["Check understanding",
         "quick whole-class checks after the Simple Rule"],
        ["Answer key and teaching notes",
         "complete Practice A/B answers (identical to the Student Book) with notes on likely confusion"],
        ["Model exchange and model response",
         "examples for rehearsal and Real Talk tasks — show these only after first attempts"],
        ["Differentiation, common errors, shorter session, review and feedback",
         "pathways, correction strategies, a 45-minute plan, and how to close the lesson"],
    ], widths=[2.6, 3.7], font_size=9.5)
    para(doc, "Back matter", size=12, bold=True, color=NAVY, space_before=6)
    para(doc, "After the sixteen guides you will find administration and key notes for the "
              "Diagnostic Test, the three Review Tests and the Final Assessment, including a model "
              "final speaking response and rubric-use reminders.")
    para(doc, "Three working rules", size=12, bold=True, color=NAVY, space_before=6)
    box(doc, "Before you teach", [
        "•  Read the Student Book lesson first — you will teach from the same page order as learners.",
        "•  Decide in advance which learners need the Intermediate scaffold and which are ready for "
        "the Advanced upgrade; both pathways share one objective.",
        "•  During speaking work, keep feedback until the two minutes are over — note errors, then "
        "repair them together afterwards.",
    ], fill=FILL_BOX_TEAL, label_color=TEAL)
    page_break(doc)


def contents(doc):
    h_banner(doc, "CONTENTS", sub="Lesson guides, front matter and assessment keys")
    para(doc, "This is an automatic Word table of contents. It updates to show page numbers when "
              "opened in Word; if it has not refreshed, select it and choose Update Field / Update "
              "entire table.", size=10.5, italic=True, color=GREY)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = 4
    p.paragraph_format.space_after = 4
    begin_run = p.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    begin_run._r.append(fld_begin)
    instr_run = p.add_run()
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = ' TOC \\o "1-1" \\h \\z \\u '
    instr_run._r.append(instr)
    sep_run = p.add_run()
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    sep_run._r.append(separate)
    result = p.add_run("Update this field to populate the contents and page numbers.")
    result.font.italic = True
    end_run = p.add_run()
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    end_run._r.append(fld_end)
    page_break(doc)


def course_overview(doc, lessons, lesson_titles):
    h_banner(doc, "COURSE OVERVIEW", sub="The approved sixteen-lesson syllabus, as taught")
    para(doc, "The course is the Student Book syllabus in full: sixteen lessons in eight thematic "
              "parts, each lesson carrying the same eight teaching functions. Nothing is merged, "
              "omitted or reordered between the two volumes.")
    rows = []
    for L in lessons:
        rows.append([str(L["n"]), lesson_titles[L["n"]], f"Part {L['part_no']}", L["goal"]])
    table(doc, ["#", "Lesson", "Part", "Speaking goal"], rows,
          widths=[0.4, 2.1, 0.7, 3.1], font_size=9)
    para(doc, "Assessment architecture", size=12, bold=True, color=NAVY, space_before=6)
    _bullets = [
        "Diagnostic Test (25 written items + matched baseline speaking sample) — before Lesson 1.",
        "Review Test 1 after Lesson 4 (scope: Lessons 1–4); Review Test 2 after Lesson 8 (1–8); "
        "Review Test 3 after Lesson 12 (1–12). Cumulative, taught material only.",
        "Final Assessment after Lesson 16: 30 written items with an item-level coverage matrix, "
        "plus the matched speaking task and four-criterion rubric.",
        "Five appendices in the Student Book: Formula Book, Common Errors Guide, Sentence "
        "Patterns, Can-Do Tracker and the Complete Answer Key (Appendix 5).",
    ]
    for b in _bullets:
        para(doc, f"•  {b}", indent=0.12)
    page_break(doc)


def teaching_principles(doc):
    h_banner(doc, "TEACHING PRINCIPLES", sub="The instructional philosophy behind every guide")
    principles = [
        ("Meaning first, form second.",
         "Every lesson opens with the Speaking Hook: learners speak before they see the rule, so "
         "the grammar arrives as a tool they were already reaching for."),
        ("Keep explanation short, production long.",
         "The Simple Rule and See the Difference should take a quarter of the lesson at most; the "
         "remaining time belongs to guided and spoken use."),
        ("Accuracy and fluency alternate.",
         "Controlled practice builds accuracy; Pair Speaking rehearses with support; the Real Talk "
         "Challenge demands independent transfer. Do not collapse the three into one."),
        ("Correction that protects speaking.",
         "Do not interrupt fluency work to fix errors. Note them, finish the task, then repair "
         "together — see the Pacing and Feedback section."),
        ("Models are examples, never scripts.",
         "Printed models illustrate effective language. Learners produce their own versions; "
           "accept any response that meets the success criteria."),
        ("One objective, two pathways.",
         "Intermediate learners secure the core pattern with scaffolds; Advanced learners deepen "
         "choice, register and precision. Advanced work must add precision, not merely workload."),
        ("Assessment stays honest.",
         "Accept legitimate alternatives, assess speaking with criteria rather than a single "
         "answer, and never present raw diagnostic and final totals as a gain score."),
        ("Return-to-rule on failure.",
         "When a self-check box is unticked, the learner goes back to the Simple Rule, redoes the "
         "matching practice item, and revisits the Formula Book card before moving on."),
    ]
    for title, text in principles:
        box(doc, title, [text], fill=FILL_BOX_GREY, label_color=GOLD)
    page_break(doc)


def pacing_and_feedback(doc):
    h_banner(doc, "PACING AND FEEDBACK", sub="Sixty minutes, mapped from the Student Book guidance")
    para(doc, "The Student Book's summary pacing table for a sixty-minute lesson groups the eight "
              "lesson functions into five blocks. The per-lesson timing tables in this book are "
              "finer cuts of exactly the same hour — every guide totals 60 minutes and every block "
              "matches the Student Book's figures.")
    table(doc, ["Student Book block", "Minutes", "Lesson functions in this book"], [
        ["Explanation and examples", "15", "Speaking Hook + Simple Rule + See the Difference"],
        ["Controlled practice", "10", "Practice the Pattern"],
        ["Pair speaking", "20", "Pair Speaking + Advanced Upgrade"],
        ["Independent speaking challenge", "10", "Real Talk Challenge"],
        ["Feedback and review", "5", "Review and Self-Check"],
        [("Total", True), ("60", True), ("", False)],
    ], widths=[2.3, 0.8, 3.2], font_size=9.5)
    para(doc, "Session variants", size=12, bold=True, color=NAVY, space_before=6)
    for b in [
        "45-minute lesson: follow the shorter-session box printed in each guide — the core "
        "objective, controlled practice and a shortened supported speaking task survive; the "
        "Advanced upgrade moves to independent study.",
        "90-minute lesson: run the full guide, then repeat the Real Talk Challenge with a new "
        "partner and finish with the Review and Self-Check written twice (second attempt in "
        "different words).",
        "Multiple weekly sessions: place a review-test day between lesson blocks rather than "
        "compressing two lessons into one hour.",
    ]:
        para(doc, f"•  {b}", indent=0.12)
    para(doc, "Feedback that protects fluency", size=12, bold=True, color=NAVY, space_before=6)
    box(doc, "During speaking tasks — do not interrupt", [
        "•  Listen with the printed success criteria; tick what you hear instead of writing while "
        "learners speak.",
        "•  Note one or two recurring global errors (tense choice, word order) on the board "
        "anonymously — never one note per learner.",
        "•  Praise one specific success per speaker: a structure, a repair, a follow-up question.",
    ], fill=FILL_BOX_TEAL, label_color=TEAL)
    box(doc, "After the task — repair in three minutes", [
        "•  Recast once, ask the class to rebuild the sentence, then say the model version.",
        "•  Correct only what was taught in this lesson; file other errors for a later review.",
        "•  Send learners back to the printed Watch out list in Part 8 for self-correction before "
        "you supply answers.",
        "•  With Advanced learners, ask them to justify their choice of structure — discussion is "
        "correction too.",
    ], fill=FILL_BOX_GREY, label_color=GOLD)
    page_break(doc)


def assessment_guidance(doc, keys_registry, model_final_speech):
    h_banner(doc, "ASSESSMENT GUIDANCE", sub="Keys, alternatives, speaking tasks and score honesty")
    para(doc, "Where the answers live", size=12, bold=True, color=NAVY)
    for b in [
        "Lesson practice answers: printed inside every guide (identical to Student Book Appendix 5).",
        "Diagnostic Test, Review Tests 1–3 and Final Assessment keys: printed in the back matter "
        "of this book; learners see them only in Student Book Appendix 5, after they have worked.",
        "Open-ended speaking tasks: never keyed to a single response. Use the success criteria in "
        "each task and the four-criterion rubric (Grammatical accuracy · Appropriate structure "
        "selection · Clarity · Successful communication; 1–4 each, total 16).",
    ]:
        para(doc, f"•  {b}", indent=0.12)
    para(doc, "Marking rules", size=12, bold=True, color=NAVY, space_before=6)
    for b in [
        "Accept every alternative the key marks as acceptable; multiple defensible answers are "
        "acknowledged in the key, not penalised.",
        "Where an item admits several readings and the key does not acknowledge them, treat it as "
        "a discussion item — do not invent a single mandatory answer.",
        "For transformations, credit any version that preserves the target structure and meaning.",
        "Record speaking scores by criterion, not as one global impression, and keep the baseline "
        "and final rubric scores side by side.",
        "Do not compare Diagnostic and Final raw grammar totals as a measure of gain — the tests "
        "differ in length and blueprint. Compare rubric criterion scores instead.",
    ]:
        para(doc, f"•  {b}", indent=0.12)
    para(doc, "Model final speaking response (one strong example, not a script)", size=12,
         bold=True, color=NAVY, space_before=6)
    box(doc, "Approximately two minutes, delivered without reading", [model_final_speech],
        fill=FILL_BOX_GOLD, label_color=GOLD)
    para(doc, "Score it with the printed rubric: note the criterion scores, then discuss with the "
              "speaker which one or two behaviours would move a 3 to a 4. Learner responses that "
              "differ in content but meet the criteria must score the same as the model.",
         size=10.5, italic=True, color=GREY)
    para(doc, "Keys — Diagnostic Test and Review Tests", size=12, bold=True, color=NAVY,
         space_before=8)
    for name in ["Diagnostic", "Review Test 1", "Review Test 2", "Review Test 3"]:
        ans = keys_registry.get(name)
        if not ans:
            continue
        para(doc, name, size=11, bold=True, color=TEAL, space_before=4)
        lines = []
        for part in ("A", "B", "C"):
            if ans.get(part):
                lines.append(f"{part}.  " + "  ·  ".join(f"{i + 1} {a}" for i, a in enumerate(ans[part])))
        box(doc, None, lines, fill=FILL_BOX_GREY, label_color=NAVY)
    para(doc, "Key — Final Assessment (30 items)", size=12, bold=True, color=NAVY, space_before=8)
    ans = keys_registry.get("Final Assessment")
    if ans:
        lines = []
        for part in ("A", "B", "C"):
            if ans.get(part):
                lines.append(f"{part}.  " + "  ·  ".join(f"{i + 1} {a}" for i, a in enumerate(ans[part])))
        box(doc, None, lines, fill=FILL_BOX_GREY, label_color=NAVY)
    para(doc, "Answer notes for the tests", size=12, bold=True, color=NAVY, space_before=6)
    for b in [
        "Review Test 1 A4 'is studying': 'these days' frames a temporary current situation; "
        "past simple here would contradict the time cue.",
        "Review Test 1 B1 'must': legal obligation; B4 'don't have to / needn't': absence of "
        "obligation — not 'mustn't', which would mean prohibition.",
        "Review Test 2 B1: accept both 'Despite the exam being hard, everyone passed' and "
        "'Despite the difficulty, everyone passed' — the first keeps the -ing complement attached "
        "to the exam, not to everyone.",
        "Review Test 3 A3: 'is said to speak' — reporting structure from Lesson 9; keep the "
        "infinitive, not a finite tense.",
        "Final Assessment A1 accepts 'have lived / 've lived / have been living' — all three "
        "describe a situation continuing to now with 'since 2020'.",
        "Final Assessment C3 accepts 'may well / probably / is likely to' and any equivalent "
        "well-formed hedge; C4 is credited with any far/much/a lot + comparative + reason frame.",
        "Diagnostic items accept the alternatives listed in the key (for example 'has worked / "
        "has been working', 'who / that', 'may / might / could', 'had / got').",
    ]:
        para(doc, f"•  {b}", indent=0.12)
    page_break(doc)
