#!/usr/bin/env python3
"""Read-only, reproducible structural audit of the independent Master Speaking.docx.

This file is NOT part of the 16-lesson Grammar for Special Speaking build.
The audit reports observable facts only; it is not a linguistic, rights or print proof.
"""
from __future__ import annotations

import json
import re
import zipfile
from collections import Counter
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "Master Speaking.docx"


def main() -> int:
    if not SOURCE.is_file():
        print("FAIL: independent manuscript is missing:", SOURCE)
        return 1
    with zipfile.ZipFile(SOURCE) as z:
        bad = z.testzip()
        if bad:
            print("FAIL: corrupt ZIP package member:", bad)
            return 1
        zip_entries = len(z.namelist())

    doc = Document(SOURCE)
    paras = list(doc.paragraphs)
    for tbl in doc.tables:
        for row in tbl.rows:
            for cell in row.cells:
                paras.extend(cell.paragraphs)
    text = "\n".join(p.text for p in paras)
    meaningful = [p for p in paras if p.text.strip()]
    words = re.findall(r"\b[\w'-]+\b", text)
    headings = [(p.style.name, p.text.strip()) for p in paras
                if p.style is not None
                and "heading" in p.style.name.lower() and p.text.strip()]
    visual_heads = [(p.style.name, p.text.strip()) for p in paras
                    if p.text.strip() and 5 <= len(p.text.strip()) <= 90
                    and any(run.bold for run in p.runs)]
    used_styles = Counter(p.style.name for p in paras if p.style is not None)
    shapes = list(doc.element.body.iter(qn("wp:docPr")))
    missing_alt = sum(not (x.get("descr") or x.get("title")) for x in shapes)
    a4_sections = sum(abs(sec.page_width.inches - 8.27) < 0.03
                      and abs(sec.page_height.inches - 11.69) < 0.03
                      for sec in doc.sections)
    tracked = sum(1 for n in doc.element.body.iter()
                  if n.tag in (qn("w:ins"), qn("w:del")))
    patterns = {
        "beginner": r"\bbeginner(s)?\b",
        "intermediate": r"\bintermediate\b",
        "advanced": r"\badvanced\b",
        "week_or_day_labels": r"\b(?:week|day)\s*\d+\b",
        "model_response": r"\bmodel response\b",
        "pair_work": r"\b(?:pair work|pair speaking|work in pairs)\b",
    }
    result = {
        "file": SOURCE.name,
        "file_bytes": SOURCE.stat().st_size,
        "docx_zip_members": zip_entries,
        "document_title": doc.core_properties.title or "",
        "document_author": doc.core_properties.author or "",
        "paragraphs_including_table_cells": len(paras),
        "nonempty_paragraphs": len(meaningful),
        "estimated_word_count": len(words),
        "tables": len(doc.tables),
        "sections": len(doc.sections),
        "a4_sections": a4_sections,
        "inline_shapes": len(doc.inline_shapes),
        "graphic_properties": len(shapes),
        "graphics_without_alt": missing_alt,
        "tracked_change_elements": tracked,
        "manual_page_breaks": sum(1 for n in doc.element.body.iter(qn("w:br"))
                                  if n.get(qn("w:type")) == "page"),
        "semantic_heading_count": len(headings),
        "heading_examples": headings[:40],
        "bold_heading_candidates": visual_heads[:30] if not headings else [],
        "most_used_styles": used_styles.most_common(12),
        "term_counts_for_editorial_review": {
            name: len(re.findall(pattern, text, re.IGNORECASE))
            for name, pattern in patterns.items()
        },
    }
    print("MASTER_SPEAKING_STRUCTURAL_AUDIT=" + json.dumps(result, ensure_ascii=False))
    print("LIMITATIONS: structural facts only; not a full manuscript reading, "
          "copyright clearance, expert pedagogical review or rendered print proof.")
    return 0 if words else 1


if __name__ == "__main__":
    raise SystemExit(main())
