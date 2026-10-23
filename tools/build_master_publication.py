#!/usr/bin/env python3
"""Reproducible PROOF exports of the independent Master Speaking.docx.

The original is never overwritten. These outputs require editorial, image-rights,
screen-reader, EPUB navigation, and visual print inspection before final release.
"""
from __future__ import annotations

import json
import re
import shutil
import subprocess
import zipfile
from pathlib import Path

import pymupdf
from docx import Document
from ebooklib import epub
from lxml import etree
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "Master Speaking.docx"
PROOF = ROOT / "proof"
TITLE = "Master Speaking (TOEFL House)"
WORKING = PROOF / f"{TITLE}.docx"
PDF = PROOF / f"{TITLE}.pdf"
EPUB = PROOF / f"{TITLE}.epub"
HTML = PROOF / f"{TITLE}.html"
REPORT = ROOT / "docs" / "master-speaking-publication-audit.md"
CSS = ROOT / "build_master" / "book.css"


def run(*args):
    print("COMMAND:", " ".join(map(str, args)), flush=True)
    subprocess.run([str(x) for x in args], cwd=ROOT, check=True, timeout=900)


def words(s):
    return len(re.findall(r"\b[\w'-]+\b", s))


def main():
    if not SOURCE.is_file():
        raise FileNotFoundError(SOURCE)
    for command in ("libreoffice", "pandoc"):
        if not shutil.which(command):
            raise RuntimeError(f"Required publisher {command} is not installed")
    PROOF.mkdir(exist_ok=True)

    original = Document(SOURCE)
    source_text = "\n".join(p.text for p in original.paragraphs)
    source_text += "\n" + "\n".join(c.text for t in original.tables for r in t.rows for c in r.cells)
    source_words = words(source_text)
    source_media = 0
    with zipfile.ZipFile(SOURCE) as zin:
        source_media = sum(p.startswith("word/media/") for p in zin.namelist())

    # Update safe, factual metadata; retain the entire original manuscript's
    # paragraphs, embedded images, tables, sections, and typography.
    original.core_properties.title = "Master Speaking"
    original.core_properties.author = "TOEFL HOUSE"
    original.core_properties.subject = "English speaking coursebook"
    original.core_properties.keywords = "speaking, communication, TOEFL HOUSE"
    original.save(WORKING)

    profile = "file:///tmp/libreoffice-master-speaking-proof"
    run("libreoffice", f"-env:UserInstallation={profile}", "--headless",
        "--convert-to", "pdf:writer_pdf_Export", "--outdir", PROOF, WORKING)
    if not PDF.is_file() or PDF.stat().st_size == 0:
        raise RuntimeError("LibreOffice failed to produce the proof PDF")

    run("pandoc", WORKING, "--from=docx", "--to=epub3", "--standalone",
        "--toc", "--toc-depth=1", "--css", CSS,
        "--metadata=title:Master Speaking", "--metadata=author:TOEFL HOUSE",
        "--metadata=lang:en", "-o", EPUB)
    run("pandoc", WORKING, "--from=docx", "--to=html5", "--standalone",
        "--self-contained", "--metadata=title:Master Speaking",
        "--metadata=lang:en", "-o", HTML)

    with pymupdf.open(PDF) as pdf:
        pdf_pages = pdf.page_count
        pdf_text = "\n".join(page.get_text() for page in pdf)
        a4_pages = sum(abs(p.rect.width - 595.28) < 3 and
                       abs(p.rect.height - 841.89) < 3 for p in pdf)
        pdf_bookmarks = len(pdf.get_toc())
    pdf_words = words(pdf_text)

    with zipfile.ZipFile(EPUB) as archive:
        names = archive.namelist()
        if names[0] != "mimetype" or archive.read("mimetype") != b"application/epub+zip":
            raise AssertionError("EPUB's first ZIP member must be the uncompressed EPUB MIME file")
        parser = etree.XMLParser(resolve_entities=False, no_network=True)
        container = etree.fromstring(archive.read("META-INF/container.xml"), parser)
        opf_name = container.xpath("string(//*[local-name()='rootfile']/@full-path)")
        if opf_name not in names:
            raise AssertionError("EPUB has no package file")
        package = etree.fromstring(archive.read(opf_name), parser)
        for item in package.xpath("//*[local-name()='manifest']/*[local-name()='item']"):
            target = str(Path(opf_name).parent / item.get("href"))
            if target not in names:
                raise AssertionError(f"EPUB manifest resource not found: {target}")
        if not package.xpath("//*[local-name()='spine']/*[local-name()='itemref']"):
            raise AssertionError("EPUB spine is empty")
        xhtml_files = [x for x in names if x.endswith(".xhtml") and "nav." not in x]
        eptext = []
        for name in xhtml_files:
            elem = etree.fromstring(archive.read(name), parser)
            eptext.extend(elem.itertext())
        epub_words = words(" ".join(eptext))

    ebook = epub.read_epub(str(EPUB), options={"ignore_ncx": False})
    epub_images = sum(1 for item in ebook.get_items()
                      if item.media_type.startswith("image/"))
    if source_words < 5000 or pdf_pages < 20 or epub_words < source_words * 0.60:
        raise AssertionError("Publication proof is incomplete: unexpected text or page loss")
    if a4_pages != pdf_pages:
        raise AssertionError(f"Unexpected non-A4 PDF pages: {pdf_pages-a4_pages}")
    if HTML.stat().st_size < 20_000:
        raise AssertionError("Generated HTML unusually small")

    report = {
        "source_words_estimate": source_words,
        "pdf_extracted_words": pdf_words,
        "epub_extracted_words": epub_words,
        "pdf_pages": pdf_pages,
        "pdf_a4_pages": a4_pages,
        "pdf_bookmarks": pdf_bookmarks,
        "source_media_files": source_media,
        "epub_embedded_images": epub_images,
        "source_graphic_properties": sum(1 for _ in original.element.body.iter(qn("wp:docPr"))),
        "proof_files": {p.name: p.stat().st_size for p in (WORKING, PDF, EPUB, HTML)},
    }
    REPORT.write_text(
        "# Master Speaking — publication proof audit\n\n"
        "> PROOF BUILDS ONLY. This report establishes structural conversion checks, "
        "not an editorial signoff, perfect image parity, accessible EPUB, third-party "
        "image rights, a qualified teaching review, or a print-ready release.\n\n"
        + "\n".join(f"- **{k.replace('_', ' ')}:** {v}" for k, v in report.items()) + "\n\n"
        "The original Master Speaking.docx is preserved unchanged. "
        "Review all pages, graphics, exercise keys, media links, navigation and EPUB "
        "rendition on real devices before approving a release.\n",
        encoding="utf-8")
    print("MASTER_PUBLICATION_REPORT=" + json.dumps(report, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
