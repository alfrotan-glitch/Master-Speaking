#!/usr/bin/env python3
"""Rebuild the original, aligned Master Speaking TEACHER'S BOOK.

Both volumes share build_master/speaking_studios.py as their lesson authority.
This manuscript includes implementable procedures, spoken models, feedback,
common corrections, assessment anchors, original Unit 1 listening script and
publisher-hosted (not copied) video resources for Units 2–8.
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.enum.style import WD_STYLE_TYPE

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/"build_master"))
sys.path.insert(0,str(ROOT/"tools"))
from speaking_studios import UNITS, COMMON_RUBRIC  # noqa
from listening_resources import LISTENING_SOURCES  # noqa

OUT=ROOT/"proof"/"Master Speaking — Teacher's Book (TOEFL House).docx"
NAVY=RGBColor(27,52,76)
TEAL=RGBColor(13,105,110)
MID=RGBColor(85,102,119)
INK=RGBColor(36,44,54)
LIGHT="EDF4F4"
LINE="DCE5E8"


def _shade(cell,fill=LIGHT):
    tcPr=cell._tc.get_or_add_tcPr()
    shd=OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def para(doc,txt="",style=None,bold=False,color=None,size=None):
    p=doc.add_paragraph(style=style)
    r=p.add_run(txt)
    if bold: r.bold=True
    if color: r.font.color.rgb=color
    if size: r.font.size=Pt(size)
    return p


def lead(doc,heading,body):
    para(doc,heading,bold=True,color=TEAL,size=10)
    para(doc,body)


def rowbox(doc,title,items):
    tbl=doc.add_table(rows=1,cols=1)
    tbl.autofit=False
    cell=tbl.cell(0,0)
    _shade(cell)
    first=cell.paragraphs[0]
    run=first.add_run(title)
    run.bold=True
    run.font.color.rgb=TEAL
    for item in items:
        cell.add_paragraph(item, style="List Bullet")
    doc.add_paragraph()


def _add_page_field(p):
    field=OxmlElement("w:fldSimple")
    field.set(qn("w:instr"),"PAGE")
    p._p.append(field)


def setup(doc):
    for sect in doc.sections:
        sect.page_width=Inches(8.27)
        sect.page_height=Inches(11.69)
        sect.left_margin=Inches(.82)
        sect.right_margin=Inches(.82)
        sect.top_margin=Inches(.82)
        sect.bottom_margin=Inches(.8)
        f=sect.footer.paragraphs[0]
        f.alignment=WD_ALIGN_PARAGRAPH.RIGHT
        f.add_run("TOEFL HOUSE  •  MASTER SPEAKING  |  TEACHER'S BOOK  •  ")
        _add_page_field(f)
    normal=doc.styles["Normal"]
    normal.font.name="Aptos"
    normal.font.size=Pt(10)
    normal.font.color.rgb=INK
    normal.paragraph_format.space_after=Pt(6)
    for name,size,color in [("Title",30,NAVY),("Heading 1",17,NAVY),("Heading 2",12,TEAL),("Heading 3",10,NAVY)]:
        st=doc.styles[name]
        st.font.name="Aptos Display" if name!="Heading 3" else "Aptos"
        st.font.size=Pt(size)
        st.font.color.rgb=color
        st.font.bold=True
        st.paragraph_format.space_before=Pt(16 if name=="Heading 1" else 10)
        st.paragraph_format.space_after=Pt(7)
        st.paragraph_format.keep_with_next=True
    for name in ["Studio Quote","Studio Note"]:
        if name not in doc.styles:
            s=doc.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH)
            s.base_style=normal
            s.font.size=Pt(9.5)
            s.font.italic=True
            s.font.color.rgb=MID
            s.paragraph_format.left_indent=Inches(.18)
    props=doc.core_properties
    props.title="Master Speaking — Teacher's Book"
    props.author="TOEFL HOUSE"
    props.subject="Unit teaching guidance and speaking assessment"
    props.keywords="communication, speaking, teacher guide, TOEFL HOUSE"


def heading(doc,title,level=1):
    return doc.add_heading(title,level)


def bullets(doc,entries):
    for item in entries:
        para(doc,item,style="List Bullet")


def cover(doc):
    for _ in range(6):doc.add_paragraph()
    p=para(doc,"TOEFL HOUSE",bold=True,color=TEAL,size=15)
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p=para(doc,"MASTER SPEAKING",bold=True,color=NAVY,size=33)
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p=para(doc,"THE TEACHER'S BOOK",bold=True,color=NAVY,size=22)
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p=para(doc,"Eight thematic units · Speaking studios · Classroom-ready tasks",color=MID,size=12)
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    for _ in range(3):doc.add_paragraph()
    p=para(doc,"Teach for interaction. Assess what learners can actually do.",color=TEAL,size=13)
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    doc.add_page_break()


def foundation(doc):
    heading(doc,"HOW TO TEACH WITH THIS BOOK")
    para(doc,"This companion is designed for one shared classroom. All learners follow the same task. Offer a brief sentence frame or preparation time where helpful; invite more independent reasoning and unscripted responses when learners are ready. Never divide activities into separate student-level tracks.")
    para(doc,"The core outcome is communicative action: students must explain a real choice, listen to another person, respond to a challenge, and revise what they say. A complicated sentence that does not communicate its intended meaning earns no special reward.")
    heading(doc,"A repeatable 60-minute studio",2)
    timing=[("Warm-up and decision",7),("Useful language in context",8),
            ("Partner role-play: first attempt",13),("Feedback and second attempt",10),
            ("Challenge discussion",8),("One-minute mini-presentation",9),
            ("Evidence, self-check and transfer",5)]
    t=doc.add_table(rows=1,cols=2)
    t.style="Light Shading Accent 1"
    t.rows[0].cells[0].text="Phase"
    t.rows[0].cells[1].text="Minutes"
    for label,value in timing:
        cells=t.add_row().cells
        cells[0].text=label
        cells[1].text=str(value)
    cells=t.add_row().cells
    cells[0].text="TOTAL"
    cells[1].text=str(sum(x[1] for x in timing))
    assert sum(x[1] for x in timing)==60
    para(doc,"If a period is shorter, protect partner speaking and revision; shorten presentation or use it as a recorded follow-up. If longer, add a second pair exchange with new partners. Avoid spending most of the lesson explaining grammar.",style="Studio Note")
    heading(doc,"Before, during, after",2)
    bullets(doc,[
        "Before: read the original unit and this studio; verify the external listening source and prepare a transcript or permitted alternative.",
        "During: give a clear mission, model a very short exchange, let pairs speak, gather evidence quietly, and return to two or three recurring problems.",
        "After: ask students what they improved on a second attempt; give one priority for the next conversation rather than a list of every mistake."
    ])
    heading(doc,"Feedback that changes the next attempt",2)
    para(doc,"Use the same sequence every time: (1) identify what the learner already did to communicate, (2) quote or paraphrase one observable problem, (3) model a more natural alternative, (4) ask the learner to repeat the same communicative move in a new turn. Do not interrupt a developing conversation for every minor error.")
    heading(doc,"A shared four-criterion speaking rubric",2)
    for title,levels in COMMON_RUBRIC:
        heading(doc,title,3)
        bullets(doc,levels)
    para(doc,"Score each dimension from 1 to 4; total /16. Use criterion-level evidence, not accent, personality, or an assumption that longer sentences are always better. When comparing performances, use the same task structure and criteria. If the prompts differ significantly, describe improvement cautiously.",style="Studio Note")
    heading(doc,"Accessible and respectful participation",2)
    para(doc,"Always allow learners to invent a character or use a fictional example. Nobody must disclose income, personal relationships, political beliefs, family conflict or health circumstances. Do not treat a person's accent as an error unless a sound prevents understanding; offer clarification and repair practice to everyone. Pair roles rotate.")
    heading(doc,"Unit map",2)
    for U in UNITS:
        para(doc,f"Unit {U['unit']} — {U['name']}: {U['can_do']}")

    heading(doc,"Optional grammar bridges across the two speaking books",2)
    para(doc,"The companion Grammar for Special Speaking course offers useful language forms, but these connections are not a prescribed calendar or separate learner pathway. Select a brief pattern only when it helps learners complete a genuine communicative task.")
    links=[
        ("Unit 1 | Listening and clarification","L3 Questions and follow-up questions; L12 Relative clauses for descriptions"),
        ("Unit 2 | Goals and contingency","L1 Tenses for real communication; L5 Conditionals for decisions"),
        ("Unit 3 | Emotions and respectful requests","L4 Modal verbs; L6 Wishes, regrets and hypothetical situations"),
        ("Unit 4 | Budget trade-offs","L7 Comparisons; L8 Linking ideas and building arguments"),
        ("Unit 5 | AI and evidence","L9 Passive voice; L15 Hedging, emphasis and natural inversion"),
        ("Unit 6 | Boundaries and agreements","L10 Causative structures; L11 Gerunds and infinitives"),
        ("Unit 7 | Global consequences","L13 Reported speech; L14 Participle and sentence patterns"),
        ("Unit 8 | Clear, responsive English","L2 Storytelling; L16 Grammar in real conversation"),
    ]
    for title,lessons in links:
        para(doc,title,bold=True,color=TEAL)
        para(doc,lessons)
    doc.add_page_break()


def unit_guide(doc,U):
    n=U["unit"]
    heading(doc,f"UNIT {n} | {U['name'].upper()}")
    para(doc,"Theme: "+U["theme"],bold=True,color=TEAL)
    rowbox(doc,"Speaking outcome",[U["mission"],U["can_do"]])
    heading(doc,"Before students speak",2)
    lead(doc,"Teaching focus",U["teacher"]["aim"])
    lead(doc,"Pre-teach in context",U["teacher"]["preteach"])
    heading(doc,"Language bank",2)
    bullets(doc,U["language"])
    heading(doc,"Opening interaction",2)
    para(doc,U["warmup"])
    heading(doc,"Guided pair-speaking mission",2)
    para(doc,U["pair"]["title"],bold=True)
    para(doc,U["pair"]["setup"])
    bullets(doc,U["pair"]["roles"])
    for i,step in enumerate(U["pair"]["steps"],1):
        para(doc,f"{i}. {step}")
    rowbox(doc,"Completion evidence",[U["pair"]["target"],U["teacher"]["evidence"]])
    heading(doc,"Short model exchange",2)
    for line in U["model"]:
        para(doc,line,style="Studio Quote")
    heading(doc,"Challenge and mini-presentation",2)
    lead(doc,"Reasoned discussion",U["challenge"])
    lead(doc,"Independent speaking",U["presentation"])
    heading(doc,"Likely error and targeted response",2)
    lead(doc,"Listen for",U["teacher"]["pitfall"])
    lead(doc,"Respond by",U["teacher"]["correction"])
    heading(doc,"Adapting support without splitting the class",2)
    lead(doc,"When the task is difficult",U["teacher"]["support"])
    lead(doc,"When the first attempt is successful",U["teacher"]["stretch"])
    heading(doc,"Assessment and next attempt",2)
    para(doc,U["feedback"])
    para(doc,"Ask learners to retry the same communicative task with a new partner, then complete a one-sentence self-check: 'Today I improved my ability to ... by ...'.")
    lead(doc,"Transfer beyond the lesson",U["transfer"])
    doc.add_page_break()


def listening(doc):
    heading(doc,"LISTENING DELIVERY AND ANSWER GUIDANCE")
    para(doc,"The Master Speaking student edition contains a resource index. It links to publicly hosted publisher pages; it does not contain licensed recordings. Preview every video locally, check captions and playback, and do not grade learners on media they cannot access. The following original Unit 1 recording script is provided for immediate classroom use.")
    heading(doc,"Unit 1 | Original active-listening passage",2)
    script=[
        "Imagine that a classmate wants to tell you about a difficult day. You put your phone away, turn toward them, and give them time to speak. You are not trying to win an argument, deliver a perfect speech, or offer advice immediately. Your first goal is to understand what happened and help your classmate feel heard.",
        "Good listening starts before the conversation. Find a place with fewer distractions. Silence notifications when possible, and prepare to pay attention. As the other person speaks, notice the main idea, not just the individual words. Do not spend the whole time planning what you will say next.",
        "Questions can help. Open-ended questions such as 'What happened next?' or 'How did that affect you?' invite fuller explanations. A yes-or-no question can be useful, but it does not always give someone much room to express an idea.",
        "Sometimes your attention drifts. That happens to everyone. Rather than pretending you understood, say, 'Sorry, I missed the last part. Could you say that again?' Asking for clarification shows that understanding matters to you.",
        "Finally, listening does not require agreement. You can listen carefully and still disagree respectfully. When people feel safe to explain themselves without being interrupted or judged immediately, they are more willing to share honest thoughts. Over time, that kind of listening can build trust and stronger relationships."
    ]
    para(doc,"Read the passage once for its main idea, and once for specific answers. Do not distribute the passage until after the first listening attempt.",style="Studio Note")
    for p in script:para(doc,p)
    heading(doc,"Unit 1 | Answer key",3)
    para(doc,"Multiple choice: 1 B (feel understood), 2 A (remove distractions), 3 B (open-ended questions), 4 B (request repetition), 5 B (trust and stronger relationships).")
    para(doc,"True/False: 1 True. 2 False — interruption does not show interest. 3 True. 4 False — listening does not require agreement. 5 True.")
    heading(doc,"Units 2–8 | Original-publisher video index",2)
    for n,theme,title,url in LISTENING_SOURCES[1:]:
        heading(doc,f"Unit {n}: {theme}",3)
        para(doc,title,bold=True)
        para(doc,url,color=TEAL,size=9)
        para(doc,"Before assessment: confirm every multiple-choice and True/False key against the actual playable recording. The teacher should have a legal transcript or suitable alternative available. No unverified key is asserted here.",style="Studio Note")
    heading(doc,"When the source cannot be played",2)
    para(doc,"Replace a viewing-dependent score with an original teacher-read passage and newly aligned questions. Keep the same speaking goal, but do not silently substitute unrelated media and retain the old answer key.")
    doc.add_page_break()


def close(doc):
    heading(doc,"FINAL PERFORMANCE AND TEACHER RECORD")
    para(doc,"Use the Unit 8 capstone: a 90-second argument linked to a unit topic, a relevant example, an acknowledged alternative view, and two unprepared follow-up questions. Ask learners to clarify or repair once if necessary.")
    heading(doc,"Two-minute feedback record",2)
    for label in ("Learner / date / prompt", "Message and reasoning /4",
                  "Listening and interaction /4", "Language choices /4",
                  "Delivery and repair /4", "One observed strength",
                  "One concrete revision", "Learner's next speaking action"):
        para(doc,label+": __________________________________________")
    para(doc,"A learner's strongest performance is not necessarily the longest or most grammatical-sounding one. Listen for a delivered message and a response to another person. Compare multiple performances before making high-stakes claims.",style="Studio Note")
    heading(doc,"Publication and safeguarding note",2)
    para(doc,"This book is a reproducible teaching resource, not proof of independent classroom effectiveness or third-party video redistribution permission. Review all practice items against the original published source, perform a print/device proof and record professional signoff before commercial release.")
    para(doc,"TOEFL HOUSE • Master Speaking • Instructor edition",bold=True,color=TEAL)


def main():
    assert len(UNITS)==8 and [u["unit"] for u in UNITS]==list(range(1,9))
    doc=Document()
    setup(doc)
    cover(doc)
    foundation(doc)
    for U in UNITS: unit_guide(doc,U)
    listening(doc)
    close(doc)
    OUT.parent.mkdir(exist_ok=True)
    doc.save(OUT)
    print("MASTER_TEACHER_BOOK:",OUT,OUT.stat().st_size, "bytes")
    return 0

if __name__=="__main__":
    raise SystemExit(main())
