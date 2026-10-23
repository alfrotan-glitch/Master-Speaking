#!/usr/bin/env python3
"""Generate, validate, and report Master Speaking's authored Teacher's Book."""
from __future__ import annotations
import json
import re
import subprocess
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET
from lxml import etree

import pymupdf
from docx import Document
import sys

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/"build_master"))
from offline_listening import OFFLINE
PROOF=ROOT/"proof"
TITLE="Master Speaking — Teacher's Book (TOEFL House)"
DOCX=PROOF/(TITLE+".docx")
PDF=PROOF/(TITLE+".pdf")
EPUB=PROOF/(TITLE+".epub")
HTML=PROOF/(TITLE+".html")
REPORT=ROOT/"docs"/"master-teacher-publication-audit.md"
CSS=ROOT/"build_master"/"book.css"


def run(*args):
    print("RUN", " ".join(map(str,args)), flush=True)
    subprocess.run([str(arg) for arg in args],cwd=ROOT,timeout=900,check=True)



def top_heading(label):
    label=" ".join(label.upper().split())
    return (label.startswith("HOW TO TEACH WITH THIS BOOK")
            or bool(re.match(r"^UNIT [1-8]\s*[|]",label))
            or label.startswith("LISTENING DELIVERY AND ANSWER GUIDANCE")
            or label.startswith("FINAL PERFORMANCE AND TEACHER RECORD"))


def curate_navigation():
    with pymupdf.open(PDF) as pdf:
        topics=[[1,title,page] for level,title,page in pdf.get_toc() if top_heading(title)]
        if len(topics)<10:
            raise AssertionError(f"Teacher PDF topic navigation incomplete: {topics}")
        pdf.set_toc(topics)
        temp=PDF.with_name(PDF.stem+".navfix.pdf")
        pdf.save(temp,garbage=4,deflate=True)
    temp.replace(PDF)
    with zipfile.ZipFile(EPUB) as z:
        names=z.namelist()
        navfile=next((name for name in names if name.endswith("nav.xhtml")),None)
        if not navfile:
            raise AssertionError("Teacher EPUB navigation file missing")
        parser=etree.XMLParser(resolve_entities=False,no_network=True)
        root=etree.fromstring(z.read(navfile),parser)
        navnodes=root.xpath("//*[local-name()='nav']")
        nav=next((n for n in navnodes if n.get("{http://www.idpf.org/2007/ops}type")=="toc"),None)
        if nav is None:
            nav=navnodes[0] if navnodes else None
        if nav is None:
            raise AssertionError("Teacher EPUB navigation element missing")
        links=[]
        used=set()
        for a in nav.xpath(".//*[local-name()='a']"):
            title=" ".join(" ".join(a.itertext()).split())
            href=a.get("href")
            if top_heading(title) and href and title not in used:
                links.append((title,href))
                used.add(title)
        if len(links)<10:
            raise AssertionError(f"Teacher EPUB topic links incomplete: {links}")
        ols=nav.xpath(".//*[local-name()='ol']")
        if not ols:
            raise AssertionError("Teacher EPUB navigation ordered list missing")
        ol=ols[0]
        for child in list(ol): ol.remove(child)
        xh="{http://www.w3.org/1999/xhtml}"
        for title,href in links:
            li=etree.SubElement(ol,xh+"li")
            etree.SubElement(li,xh+"a",href=href).text=title
        edited={navfile:etree.tostring(root,encoding="utf-8",xml_declaration=True)}
        ncx=next((name for name in names if name.endswith(".ncx")),None)
        if ncx:
            tree=etree.fromstring(z.read(ncx),parser)
            maps=tree.xpath("//*[local-name()='navMap']")
            if maps:
                mp=maps[0]
                for child in list(mp): mp.remove(child)
                ns="{http://www.daisy.org/z3986/2005/ncx/}"
                for i,(title,href) in enumerate(links,1):
                    point=etree.SubElement(mp,ns+"navPoint",id=f"topic-{i}",playOrder=str(i))
                    label=etree.SubElement(point,ns+"navLabel")
                    etree.SubElement(label,ns+"text").text=title
                    etree.SubElement(point,ns+"content",src=href)
                edited[ncx]=etree.tostring(tree,encoding="utf-8",xml_declaration=True)
        temp=EPUB.with_name(EPUB.stem+".navfix.epub")
        with zipfile.ZipFile(temp,"w") as out:
            for info in z.infolist():
                out.writestr(info,edited.get(info.filename,z.read(info.filename)))
    temp.replace(EPUB)
    return len(topics),len(links)


def validate():
    for file in [DOCX,PDF,EPUB,HTML]:
        if not file.is_file() or file.stat().st_size<10_000:
            raise AssertionError(f"Missing/truncated teacher artifact: {file}")
    doc=Document(DOCX)
    text="\n".join(p.text for p in doc.paragraphs)
    headings=[p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
    for n in range(1,9):
        if not any(h.startswith(f"UNIT {n} |") for h in headings):
            raise AssertionError(f"Teacher guide missing Unit {n}")
    for n,passage in OFFLINE.items():
        if passage["title"] not in text:
            raise AssertionError(f"Teacher guide missing Unit {n} offline listening passage")
        for question,answer in passage["checks"]:
            if question not in text or answer not in text:
                raise AssertionError(f"Teacher guide missing listening question/answer in unit {n}")
    for marker in ["LISTENING DELIVERY AND ANSWER GUIDANCE",
                   "FINAL PERFORMANCE AND TEACHER RECORD",
                   "HOW TO TEACH WITH THIS BOOK",
                   "A shared four-criterion speaking rubric"]:
        if marker not in text:
            raise AssertionError(f"Missing teacher section {marker}")
    if not all(abs(s.page_width.inches-8.27)<.03 and
               abs(s.page_height.inches-11.69)<.03 for s in doc.sections):
        raise AssertionError("Teacher DOCX is not A4")
    with pymupdf.open(PDF) as pdf:
        count=len(pdf)
        contents="\n".join(p.get_text() for p in pdf)
        if count<10 or count>85:
            raise AssertionError(f"Unexpected teacher page count {count}")
        if not all(abs(p.rect.width-595.28)<3 and abs(p.rect.height-841.89)<3 for p in pdf):
            raise AssertionError("Teacher PDF pages must all be A4")
        if not all(f"UNIT {n}" in contents for n in range(1,9)):
            raise AssertionError("Teacher PDF lost one or more unit guides")
        if len(pdf.get_toc())<8:
            raise AssertionError("Teacher PDF lacks useful bookmarks")
        bookmarks=len(pdf.get_toc())
    with zipfile.ZipFile(EPUB) as z:
        names=z.namelist()
        if names[0]!="mimetype" or z.read("mimetype")!=b"application/epub+zip":
            raise AssertionError("EPUB MIME/first entry invalid")
        if "META-INF/container.xml" not in names:
            raise AssertionError("EPUB missing container")
        for name in names:
            if name.endswith((".xhtml",".opf",".ncx")):
                ET.fromstring(z.read(name))
        spine="\n".join(z.read(name).decode("utf-8",errors="replace")
                        for name in names if name.endswith(".xhtml"))
        for n in range(1,9):
            if f"UNIT {n}" not in spine:
                raise AssertionError(f"Teacher EPUB lost Unit {n}")
    if HTML.stat().st_size<30_000:
        raise AssertionError("Teacher HTML unexpectedly small")
    details={
        "unit_count":8,
        "teacher_pdf_pages":count,
        "teacher_pdf_bookmarks":bookmarks,
        "teacher_docx_headings":len(headings),
        "teacher_paragraphs":len(doc.paragraphs),
        "teacher_files":{p.name:p.stat().st_size for p in [DOCX,PDF,EPUB,HTML]},
        "scope":"Technical edition checks, not signed pedagogical/rights/visual approval",
    }
    REPORT.write_text("# Master Speaking — Teacher edition proof audit\n\n"+
                      "\n".join(f"- **{key.replace('_',' ')}:** {value}" for key,value in details.items())+
                      "\n\nThe eight original teacher guides share the student studio data source; original Unit 1 audio is teacher-authored. Third-party media is linked, not copied. Qualitative editing and physical proof still require review.\n",encoding="utf-8")
    print("TEACHER_PROOF_AUDIT="+json.dumps(details,ensure_ascii=False))
    return details


def main():
    run("python","tools/build_master_teacher.py")
    run("libreoffice","-env:UserInstallation=file:///tmp/libreoffice-master-teacher",
        "--headless","--convert-to","pdf:writer_pdf_Export","--outdir",PROOF,DOCX)
    run("pandoc",DOCX,"--from=docx","--to=epub3","--standalone",
        "--toc","--toc-depth=2","--css",CSS,
        "--metadata=title:Master Speaking — Teacher's Book",
        "--metadata=author:TOEFL HOUSE","--metadata=lang:en","-o",EPUB)
    run("pandoc",DOCX,"--from=docx","--to=html5","--standalone",
        "--self-contained","--metadata=lang:en","-o",HTML)
    pdf_topics,epub_topics=curate_navigation()
    print(f"TEACHER_NAVIGATION: PDF {pdf_topics}, EPUB {epub_topics}")
    validate()
    return 0


if __name__=="__main__":
    raise SystemExit(main())
