# -*- coding: utf-8 -*-
"""Final de-scheduling fixes (substring-based). Grammar-content words stay."""
import re
from docx import Document

PATH = '/home/user/Master-Speaking/Master Grammar C1-C2 (TOEFL House).docx'

RUN_REPLACES = [
    ("PART ONE — MONTH 1", "PART ONE"),
    ("PART TWO — MONTH 2", "PART TWO"),
    ("— FINAL MASTERY DAY", "— FINAL MASTERY"),
    ("— Final Mastery Day", "— Final Mastery"),
    ("Homework & Course Self-Check", "Practice & Self-Check"),
    ("→  TOEFL Task  →  Homework  →", "→  TOEFL Task  →  Practice  →"),
    ("→  TOEFL Task  →  HOMEWORK  →", "→  TOEFL Task  →  PRACTICE  →"),
    ("Every lesson ends with homework and a self-check",
     "Every unit ends with practice and a self-check"),
    ("Every homework is a speaking recording. Listen once before sending it to your teacher.",
     "Every recording task is real speaking practice."),
    ("After each lesson day, review the Formula Card here:",
     "After each unit, review the Formula Card here:"),
]

# delete the 'Weekly' bullet (course scheduling)
DELETE_KEYS = ["a 2-minute recorded answer using all the week's structures",
               "Weeks 1, 3, 5, 7"]

# rewrite the teacher-tip bullet (whole paragraph)
TIP_OLD_KEY = "Some units are heavier on form and accuracy; others on fluency and production — pace yourself accordingly."
# (already rewritten last run in one case; handle the raw bullet too)
PARA_REWRITES = [
    ("•  Day A of each week",
     "•  Some units are heavier on form and accuracy; others on fluency and production — "
     "pace yourself accordingly."),
]


def fix_runs(p):
    n = 0
    for r in p.runs:
        t = r.text
        for old, new in RUN_REPLACES:
            if old in t:
                t = t.replace(old, new)
                n += 1
        r.text = t
    return n


def main():
    doc = Document(PATH)
    fixed = 0

    for p in list(doc.paragraphs):
        fixed += fix_runs(p)
        t = p.text.strip()
        # delete scheduling bullet
        if t.startswith("•") and any(k in t for k in DELETE_KEYS):
            p._p.getparent().remove(p._p)
            fixed += 1
            continue
        # rewrite day-A tip bullet
        for key, new in PARA_REWRITES:
            if t.startswith(key):
                for r in list(p.runs):
                    r._element.getparent().remove(r._element)
                run = p.add_run(new)
                run.font.name = "Calibri"
                from docx.shared import Pt
                run.font.size = Pt(11)
                fixed += 1

    for tb in doc.tables:
        for row in tb.rows:
            for c in row.cells:
                for p in c.paragraphs:
                    fixed += fix_runs(p)

    doc.save(PATH)
    print("fixes applied:", fixed)

    # report
    doc2 = Document(PATH)
    texts = [p.text for p in doc2.paragraphs]
    for t in doc2.tables:
        for r in t.rows:
            for c in r.cells:
                texts.append(c.text)
    blob = "\n".join(texts)
    print("\n=== remaining (should be grammar content only) ===")
    for w in ["Day ", "DAY ", "Week", "week", "Month", "month", "MONTH", "Homework", "homework", "teaching days"]:
        for t in texts:
            for m in re.finditer(re.escape(w), t):
                s = max(0, m.start()-35)
                print(f"  [{w}] …{t[s:m.end()+38].strip().replace(chr(10),' ')}…")
EOF_MARKER_NOT_USED = None


if __name__ == "__main__":
    main()
