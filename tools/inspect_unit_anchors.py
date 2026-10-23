#!/usr/bin/env python3
"""Identify actual unit headings and insertion points in the original manuscript."""
from pathlib import Path
import re
from docx import Document
ROOT=Path(__file__).resolve().parents[1]
SRC=ROOT/"Master Speaking.docx"
OUT=ROOT/"docs"/"master-unit-anchor-map.md"

def main():
    doc=Document(SRC)
    paragraphs=list(doc.paragraphs)
    lines=["# Master Speaking unit anchor map","","Paragraph indices are from document.paragraphs, not Word page numbers.",""]
    for i,p in enumerate(paragraphs):
        t=" ".join(p.text.split())
        if re.search(r"(?i)\bUNIT\s*[1-8]\s*[:—–-]",t) or re.match(r"(?i)^\s*CHAPTER\s+[1-5]\b",t):
            lines.extend([f"## Paragraph {i}: {t[:180]}",
                          f"Style: {p.style.name if p.style else '-'}",
                          "Neighbors:"])
            for offset in range(-2,4):
                j=i+offset
                if 0<=j<len(paragraphs):
                    q=paragraphs[j]
                    lines.append(f"- {j} [{q.style.name if q.style else '-'}]: {q.text.strip()[:150]}")
            lines.append("")
    OUT.write_text("\n".join(lines),encoding="utf-8")
    print("ANCHOR_MAP:",len(paragraphs),"paragraphs",len(lines),"report lines")
    return 0

if __name__=="__main__":
    raise SystemExit(main())
