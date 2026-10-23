#!/usr/bin/env python3
"""Twenty independent editorial and publication review lenses.

Enforce provable invariants. Report unresolved expert and accessibility work
honestly; never equate green CI with a perfect book.
"""
from __future__ import annotations

import posixpath
import re
import sys
import zipfile
from pathlib import Path
from urllib.parse import unquote, urlsplit

import pymupdf
from docx import Document
from docx.oxml.ns import qn
from lxml import etree

ROOT = Path(__file__).resolve().parents[1]
REPORT = ROOT / "docs" / "quality-audit-20-dimensions.md"
MASTER = ROOT / "Master Speaking.docx"
PROOF = ROOT / "proof" / "Master Speaking (TOEFL House).docx"
PDF = PROOF.with_suffix(".pdf")
EPUB = PROOF.with_suffix(".epub")
STUDENT = ROOT / "dist" / "Grammar for Special Speaking (TOEFL House).docx"
TEACHER = ROOT / "dist" / "Grammar for Special Speaking — Teacher's Book (TOEFL House).docx"


def text_of(doc):
    return "\n".join(p.text for p in doc.paragraphs) + "\n" + "\n".join(
        c.text for t in doc.tables for r in t.rows for c in r.cells)


def internal_epub_links(path):
    errors, checked, external = [], 0, 0
    with zipfile.ZipFile(path) as z:
        names = set(z.namelist())
        parser = etree.XMLParser(resolve_entities=False, no_network=True)
        docs = {}
        for name in names:
            if name.endswith((".xhtml", ".svg")):
                docs[name] = etree.fromstring(z.read(name), parser)
        ids = {name: set(node.xpath("//@id")) for name, node in docs.items()}
        for name, node in docs.items():
            for element in node.xpath("//*[@href or @src]"):
                for key in ("href", "src"):
                    href = element.get(key)
                    if not href or href.startswith("data:"):
                        continue
                    parsed = urlsplit(href)
                    if parsed.scheme or parsed.netloc:
                        external += 1
                        continue
                    dest = (posixpath.normpath(posixpath.join(
                        posixpath.dirname(name), unquote(parsed.path)))
                        if parsed.path else name)
                    if dest not in names:
                        errors.append(f"{name}: missing internal {key}={href}")
                    elif parsed.fragment and dest in ids and unquote(parsed.fragment) not in ids[dest]:
                        errors.append(f"{name}: broken fragment {href}")
                    checked += 1
    return errors, checked, external


def main():
    original = Document(MASTER)
    working = Document(PROOF)
    student = Document(STUDENT)
    teacher = Document(TEACHER)
    source_text, working_text = text_of(original), text_of(working)
    g_text, t_text = text_of(student), text_of(teacher)
    src_count, dst_count = len(source_text.split()), len(working_text.split())
    graphics = list(working.element.body.iter(qn("wp:docPr")))
    missing_alt = [(i+1, node.get("name", "") or node.get("id", ""))
                   for i,node in enumerate(graphics)
                   if not (node.get("descr") or node.get("title"))]
    heads1 = [p.text.strip() for p in working.paragraphs if p.style.name == "Heading 1"]
    heads2 = [p.text.strip() for p in working.paragraphs if p.style.name == "Heading 2"]
    bad_heads = [p.text.strip() for p in working.paragraphs
                 if p.style.name.startswith("Heading") and re.match(
                     r"(?i)^(?:[0-9]+[.)]\s+|_{3,}|[A-D]\.\s+|choose the best answer)", p.text.strip())]
    links_bad, links_tested, external_links = internal_epub_links(EPUB)
    with pymupdf.open(PDF) as pdf:
        pages = pdf.page_count
        a4 = sum(abs(page.rect.width-595.28)<3 and abs(page.rect.height-841.89)<3 for page in pdf)
        bookmarks = [e[1] for e in pdf.get_toc()]
        blank = [i+1 for i,p in enumerate(pdf) if len(p.get_text().strip()) < 20]
    student_ok = all(f"LESSON {i}" in g_text for i in range(1,17))
    teacher_ok = all(f"LESSON {i}" in t_text for i in range(1,17))
    teacher_guides = t_text.count("TEACHER GUIDE")
    failures = []
    if abs(dst_count-src_count)/max(1,src_count) > .02:
        failures.append(f"Word proof lost or added >2% text ({src_count} vs {dst_count} tokens)")
    if len(heads1)<10 or len(heads1)>25 or bad_heads:
        failures.append(f"Working Word headings: H1={len(heads1)}, misplaced prompts={bad_heads[:6]}")
    if pages!=a4 or blank:
        failures.append(f"PDF proof A4 pages={a4}/{pages}, blank pages={blank}")
    if len(bookmarks)!=12:
        failures.append(f"PDF navigation has {len(bookmarks)} bookmarks rather than 12")
    if links_bad:
        failures.extend(links_bad[:30])
    if not (student_ok and teacher_ok and teacher_guides>=16):
        failures.append("Student/Teacher grammar lessons or guide markers are missing")
    if working.core_properties.author != "TOEFL HOUSE":
        failures.append("Master working copy lacks expected publisher metadata")
    rows=[
        ("Source integrity & metadata","PASS" if not abs(dst_count-src_count)/max(1,src_count)>.02 else "BLOCKER",f"Original/proof tokens {src_count}/{dst_count}; publisher {working.core_properties.author}"),
        ("Audience and measurable learning objectives","REVIEW","Expert inspection and course-level learning-outcome review required"),
        ("Course architecture and unit consistency","REVIEW",f"{len(heads1)} top-level topic headings and {len(heads2)} section headings"),
        ("Skill progression and cognitive load","REVIEW","Review learning sequence and activity demand with real learners"),
        ("Grammatical accuracy and naturalness","REVIEW","Qualified full-text line-by-line language edit required"),
        ("Vocabulary, collocations and recycling","REVIEW","Check topic relevance, lexical recurrence and examples"),
        ("Speaking time and genuine interaction","REVIEW","Run an observed classroom trial measuring learner talk"),
        ("Listening/audio/media reliability","REVIEW","Verify source URLs, licenses, transcripts and accessibility"),
        ("Reading-to-speaking transfer","REVIEW","Audit reading activities for purposeful oral outputs"),
        ("Pronunciation and intelligibility","REVIEW","Expert phonetics and audio modeling review required"),
        ("Practice-item clarity and answer alternatives","REVIEW","Check every prompt, distractor and response alternative"),
        ("Assessment, rubrics, scorer reliability","REVIEW","Pilot repeatability, scoring anchors, trained raters"),
        ("Teacher guide and key completeness","PASS" if student_ok and teacher_ok and teacher_guides>=16 else "BLOCKER",f"Grammar 16 student lessons={student_ok}; teacher lessons={teacher_ok}; guide tags={teacher_guides}; quality still requires review"),
        ("Cultural inclusion and safe participation","REVIEW","Offer non-disclosure options and culturally fair topics"),
        ("Editable Word semantic headings","PASS" if len(heads1)>=10 and not bad_heads else "BLOCKER",f"Topic headings={len(heads1)}, accidental question headings={len(bad_heads)}"),
        ("Typography, layout and print proof","REVIEW",f"{a4}/{pages} A4 pages; full visual and physical proof remain unverified"),
        ("Images, alt text and accessible contrast","BLOCKER" if missing_alt else "REVIEW",f"{len(graphics)} graphic objects; {len(missing_alt)} lack alt/title; image-by-image descriptions and color contrast inspection required"),
        ("EPUB internal navigation and links","PASS" if not links_bad and len(bookmarks)==12 else "BLOCKER",f"{len(bookmarks)} curated PDF bookmarks; {links_tested} checked internal links; {len(links_bad)} broken"),
        ("Cross-format text and image fidelity","REVIEW",f"Original/proof tokens {src_count}/{dst_count}; compare images and tables in rendered copies"),
        ("Intellectual property, sources and approval","REVIEW","Rights for media, bibliography, ISBN, and independent publication signoff required"),
    ]
    verified=sum(x[1]=="PASS" for x in rows)
    blockers=sum(x[1]=="BLOCKER" for x in rows)
    review=sum(x[1]=="REVIEW" for x in rows)
    body=["# TOEFL HOUSE — 20-dimension evidence-led quality assessment",
          "",
          f"**Current results:** {verified} machine-verified, {blockers} publication blocker(s), {review} expert-review categories.",
          "",
          "A PASS establishes only the specific automated invariant; it is not a score of 100%. A REVIEW is not a failure, but the book cannot be certified flawless without evidence.",
          "",
          "| # | Lens | State | Observed evidence / outstanding test |",
          "|---:|---|---|---|"]
    for i,(title,state,evidence) in enumerate(rows,1):
        body.append(f"| {i} | {title} | **{state}** | {evidence.replace('|','/')} |")
    body += ["","## Graphics needing meaningful alt text / decorative treatment",
             f"{len(missing_alt)} elements without an alt/title description. No invented visual descriptions have been added.",
             "",
             "| Number | Name / ID |","|---:|---|"]
    for i,name in missing_alt:
        body.append(f"| {i} | {name.replace('|','/')} |")
    body += ["","## Blocking automated defects",""]
    body += ["- "+x for x in failures] if failures else ["- None found among enforced structural checks."]
    body += ["",
             "## Honest release constraints",
             "Human teaching trials, exhaustive language editing, visual pagination proof, image rights and accessibility, EPUBCheck, media verification, and print proof are not certified by this run.",
             ""]
    REPORT.write_text("\n".join(body),encoding="utf-8")
    print(f"20-LENS AUDIT: PASS={verified}, BLOCKERS={blockers}, REVIEW={review}, MACHINE_FAILURES={len(failures)}")
    print(f"EPUB checked {links_tested} internal resources and found {len(links_bad)} broken; {external_links} external URL(s).")
    for problem in failures: print("FAIL:",problem)
    return 1 if failures else 0


if __name__ == "__main__":
    sys.exit(main())
