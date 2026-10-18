#!/usr/bin/env python3
"""Structural and cross-format checks for the three final distribution files."""
from __future__ import annotations

import re
import shutil
import sys
import zipfile
from pathlib import Path
from urllib.parse import unquote, urlparse

import pymupdf as fitz
from docx import Document
from lxml import etree
from ebooklib import epub

ROOT = Path(__file__).resolve().parents[1]
TITLE = "Grammar for Special Speaking (TOEFL House)"
DIST = ROOT / "dist"
DOCX = DIST / f"{TITLE}.docx"
PDF = DIST / f"{TITLE}.pdf"
EPUB = DIST / f"{TITLE}.epub"


def norm(text):
    return re.sub(r"[^a-z0-9]+", "", text.lower())


def docx_text(path):
    d = Document(path)
    out = []
    for child in d.element.body:
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            out.append("".join(t.text or "" for t in child.iter() if t.tag.endswith("}t")))
        elif tag == "tbl":
            out.extend(t.text or "" for t in child.iter() if t.tag.endswith("}t"))
    return "\n".join(out)


def epub_validate(path):
    with zipfile.ZipFile(path) as z:
        names = z.namelist()
        assert names and names[0] == "mimetype", "mimetype is not the first ZIP entry"
        assert z.read("mimetype") == b"application/epub+zip", "incorrect EPUB mimetype"
        parser = etree.XMLParser(resolve_entities=False, no_network=True)
        docs = {}
        for name in names:
            if name.endswith((".xhtml", ".opf", ".ncx", ".xml")):
                docs[name] = etree.fromstring(z.read(name), parser)
        root = etree.fromstring(z.read("META-INF/container.xml"), parser)
        opf_path = root.xpath("string(//*[local-name()='rootfile']/@full-path)")
        assert opf_path in names, "container.xml package path is missing"
        opf = docs[opf_path]
        opf_dir = Path(opf_path).parent
        manifest = {}
        for item in opf.xpath("//*[local-name()='manifest']/*[local-name()='item']"):
            href = unquote(item.get("href"))
            target = str(opf_dir / href) if str(opf_dir) != "." else href
            target = str(Path(target))
            assert target in names, f"manifest item missing: {target}"
            manifest[item.get("id")] = (target, item.get("media-type"))
        for ref in opf.xpath("//*[local-name()='spine']/*[local-name()='itemref']"):
            assert ref.get("idref") in manifest, f"spine refers to missing id {ref.get('idref')}"
        nav_items = [p for p, m in manifest.values() if m == "application/xhtml+xml" and "nav" in p.lower()]
        assert nav_items, "EPUB has no navigation document"
        for path_name, node in docs.items():
            if not path_name.endswith((".xhtml", ".ncx")):
                continue
            base = Path(path_name).parent
            ids = set(node.xpath("//@id"))
            for link in node.xpath("//*[local-name()='a']/@href"):
                parsed = urlparse(link)
                if parsed.scheme or parsed.netloc:
                    continue
                target = str((base / unquote(parsed.path)).resolve()) if parsed.path else path_name
                # resolve against the in-package path without allowing filesystem access
                target = str(Path(path_name).parent.joinpath(unquote(parsed.path))) if parsed.path else path_name
                target = str(Path(target))
                if target not in names:
                    raise AssertionError(f"broken EPUB href {link!r} in {path_name}")
                if parsed.fragment:
                    target_node = docs.get(target)
                    assert target_node is not None and parsed.fragment in set(target_node.xpath("//@id")), \
                        f"broken EPUB fragment {link!r} in {path_name}"
        epub_doc = next((n for n, m in manifest.values() if n.endswith("book.xhtml")), None)
        assert epub_doc, "EPUB book XHTML is missing"
        text = " ".join(docs[epub_doc].itertext())
        assert text.count("LESSON ") >= 16, "EPUB is missing lessons"
        assert all(s in text for s in ("REVIEW TEST 1", "REVIEW TEST 2", "REVIEW TEST 3", "FINAL ASSESSMENT", "APPENDIX 5")), "EPUB missing a required section"
        # EbookLib's parser provides an independent package-open check.
        epub.read_epub(str(path), options={"ignore_ncx": False})
        return text, len(names)


def main():
    failures = []
    def check(label, fn):
        try:
            value = fn()
            print(f"PASS {label}" + (f" — {value}" if value is not None else ""))
        except Exception as e:
            failures.append((label, str(e)))
            print(f"FAIL {label} — {e}")

    for f in (DOCX, PDF, EPUB):
        check(f"exists: {f.relative_to(ROOT)}", lambda f=f: f"{f.stat().st_size:,} bytes" if f.is_file() and f.stat().st_size > 0 else (_ for _ in ()).throw(FileNotFoundError(f)))

    dtext = ""
    def check_docx():
        nonlocal dtext
        doc = Document(DOCX)
        dtext = docx_text(DOCX)
        assert doc.sections and all(abs(s.page_width.inches-8.27)<0.03 and abs(s.page_height.inches-11.69)<0.03 for s in doc.sections), "not A4"
        assert all(f"LESSON {i}" in dtext for i in range(1,17))
        assert all(f"APPENDIX {i}" in dtext for i in range(1,6))
        assert "FINAL ASSESSMENT" in dtext and "REVIEW TEST 3" in dtext
        headings = [p for t in doc.tables for r in t.rows for c in r.cells for p in c.paragraphs if p.style.name.startswith("Heading")]
        assert sum(p.style.name.startswith("Heading 1") for p in headings) >= 20, "semantic Heading 1 styles missing"
        assert sum(p.style.name.startswith("Heading 2") for p in headings) >= 128, "semantic lesson part headings missing"
        usable = (doc.sections[0].page_width-doc.sections[0].left_margin-doc.sections[0].right_margin)/914400
        assert all(sum(c.width/914400 for c in t.rows[0].cells) <= usable+0.06 for t in doc.tables), "a table exceeds the printable text width"
        footer_xml = " ".join(s.footer._element.xml for s in doc.sections)
        assert "PAGE" in footer_xml, "live page-number field missing"
        assert doc.core_properties.title == "Grammar for Special Speaking" and doc.core_properties.author == "TOEFL HOUSE"
        return f"editable DOCX; semantic headings {len(headings)}; A4; table widths and footer field checked"
    check("DOCX structure and required contents", check_docx)

    pdftext = ""
    def check_pdf():
        nonlocal pdftext
        p = fitz.open(PDF)
        assert p.page_count > 40, f"unexpectedly short: {p.page_count} pages"
        assert all(abs(page.rect.width-595.28)<2 and abs(page.rect.height-841.89)<2 for page in p), "PDF is not A4"
        pdftext = "\n".join(page.get_text() for page in p)
        assert len(pdftext) > 100000, "PDF has too little searchable text"
        assert "LESSON 16" in pdftext and "FINAL ASSESSMENT" in pdftext and "APPENDIX 5" in pdftext
        assert all(len(page.get_text().strip()) > 30 for page in p), "PDF has an empty/near-empty page"
        assert len(p.get_toc()) >= 24, f"PDF bookmarks/navigation missing: {len(p.get_toc())} entries"
        assert len(p[3].get_links()) >= 20, "page-numbered contents has too few internal links"
        embedded = [font for font in p.get_page_fonts(0, full=True) if "DejaVuSans" in font[3] and font[1] == "ttf"]
        assert embedded and all(p.extract_font(font[0])[3] for font in embedded), "Unicode font is not embedded"
        assert "\ufffd" not in pdftext, "replacement glyph present in extracted PDF text"
        return f"{p.page_count} A4 pages; searchable text; {len(p.get_toc())} bookmarks"
    check("PDF structure, text, pages and bookmarks", check_pdf)

    epubtext = ""
    def _epub_result():
        nonlocal epubtext
        epubtext, count = epub_validate(EPUB)
        return f"valid package structure; {count} ZIP entries; local links resolved"
    check("EPUB package, XML, links, navigation and reflowable XHTML", _epub_result)

    if dtext and pdftext and epubtext:
        nd, np, ne = norm(dtext), norm(pdftext), norm(epubtext)
        terms = [f"LESSON {i}" for i in range(1,17)] + [f"APPENDIX {i}" for i in range(1,6)] + ["DIAGNOSTIC TEST", "REVIEW TEST 1", "REVIEW TEST 2", "REVIEW TEST 3", "FINAL ASSESSMENT", "Answer key", "Speaking task"]
        missing = [(term, fmt) for term in terms for fmt, blob in (("DOCX",nd),("PDF",np),("EPUB",ne)) if norm(term) not in blob]
        check("cross-format required-section consistency", lambda: f"{len(terms)} markers in all 3 editions" if not missing else (_ for _ in ()).throw(AssertionError(str(missing[:10]))))

    if shutil.which("epubcheck"):
        import subprocess
        result = subprocess.run(["epubcheck", str(EPUB)], capture_output=True, text=True)
        print("EPUBCheck:", result.stdout.strip() or result.stderr.strip())
        if result.returncode:
            failures.append(("EPUBCheck", result.stdout + result.stderr))
    else:
        print("INFO EPUBCheck not installed; package was checked for XML well-formedness, manifest/spine, navigation and local links. This is not an EPUBCheck result.")
    if failures:
        print("\nFAILURES:")
        for name, message in failures:
            print(f"- {name}: {message}")
        return 1
    print("\nRESULT: all available structural and cross-format checks passed.")
    return 0

if __name__ == "__main__":
    sys.exit(main())
